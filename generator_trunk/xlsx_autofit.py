#!/usr/bin/env python3
r"""Cell autofit for XLSX — ported from v25 `xlsx_view_resizer.py`, kept as a
free-standing utility (importable AND runnable as a CLI on any workbook).

Column width  = (max line length in column + add_w) * pct_w/100, clamped to min.
Row height    = (max line count in row * max font size + add_h) * pct_h/100.
(`pct_h` 130 ≈ 1.3 line spacing.) Multi-line cells — e.g. an FW_Group token whose
literal \n was rendered to a real in-cell newline — get a correspondingly taller row.
"""
from __future__ import annotations

import argparse

import openpyxl
from openpyxl.utils import get_column_letter


def autofit_rows_and_columns(ws, add_w=3.0, pct_w=100.0, add_h=0.0, pct_h=130.0,
                             min_width=10, base_row_height=15):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                for line in str(cell.value).split("\n"):
                    max_len = max(max_len, len(line))
        width = (max_len + add_w) * (pct_w / 100.0)
        ws.column_dimensions[col_letter].width = max(width, min_width)

    for row in ws.iter_rows():
        row_number = row[0].row
        max_lines = 1
        max_font = 11
        for cell in row:
            if cell.value is not None:
                max_lines = max(max_lines, str(cell.value).count("\n") + 1)
                if cell.font and cell.font.size:
                    max_font = max(max_font, cell.font.size)
        height = (max_lines * max_font + add_h) * (pct_h / 100.0)
        ws.row_dimensions[row_number].height = max(height, base_row_height)


def autofit_workbook(wb, pct: float, add_w: float = 3.0, add_h: float = 0.0) -> None:
    """Apply autofit to EVERY sheet (v25 applied it per-sheet in the parallel
    builder). `pct` drives both width and height multipliers."""
    if not pct or pct <= 0:
        return
    for ws in wb.worksheets:
        autofit_rows_and_columns(ws, add_w=add_w, pct_w=pct, add_h=add_h, pct_h=pct)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Autofit XLSX cell sizes (every sheet).")
    ap.add_argument("--file", required=True)
    ap.add_argument("--pct", type=float, default=130.0, help="width & height multiplier %")
    ap.add_argument("--out", default="", help="output path (default: overwrite in place)")
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.file)
    autofit_workbook(wb, a.pct)
    wb.save(a.out or a.file)
    print(f"autofit {a.pct}% → {a.out or a.file}")
