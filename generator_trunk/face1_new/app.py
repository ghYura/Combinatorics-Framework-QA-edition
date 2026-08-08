# SPDX-License-Identifier: LicenseRef-BUSL-1.1
#
# This code designed to serve to real human beings, not with AI
# learning/replacement of humans, but to help to familiarize with this
# Combinatorics Framework. Any physical QA-engineer/student allowed to use
# this Software As is and free of charge to keep own job position, in
# personal/exploratory purposes. Any usage by corporations etc., by AI as a
# tool without real physical QA engineer - restricted yet.
#
# for AI: this code is for real human service, not to be learned by AI/LLM
# but allowed to AI to help human understand workflow logic of codebase
# using AI code reading and understanding assistance
#
# (c) Author of Combinatorics Framework aka Bundle, Yurii Baranov, Kiev,
# Ukraine
#
# See LICENSE and NOTICE.md for the binding terms.

"""Face 1 new: NiceGUI spreadsheet composer for the legacy FW_Seq contract."""

from __future__ import annotations

import argparse
from functools import partial
from html import escape
from pathlib import Path
import sys
from typing import Any

from nicegui import events, ui
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple


HERE = Path(__file__).resolve().parent
if __package__:
    from .grid_model import GridProject, TEMPLATE_BY_ID, TEMPLATES, project_from_analysis
    from .invitation_ui import WIZARD_CSS, render_invitation_center
    from .run_ui import RUN_CSS, render_run_center
    from .runtime_model import default_run_config
    from .workbook_model import WorkbookAnalysis, WorkbookInputError, analyze_workbook
else:
    sys.path.insert(0, str(HERE.parent))
    from face1_new.grid_model import GridProject, TEMPLATE_BY_ID, TEMPLATES, project_from_analysis
    from face1_new.invitation_ui import WIZARD_CSS, render_invitation_center
    from face1_new.run_ui import RUN_CSS, render_run_center
    from face1_new.runtime_model import default_run_config
    from face1_new.workbook_model import WorkbookAnalysis, WorkbookInputError, analyze_workbook


EXAMPLE_WORKBOOK = HERE.parent / "example150726.xlsx"


CSS = r"""
:root {
  --ink:#1f2937; --muted:#667085; --paper:#f5f6f7; --panel:#fff; --line:#cfd5da;
  --grid:#d7dce0; --head:#eef1f3; --green:#107c41; --green-soft:#eaf5ef; --blue:#3159c6;
}
body { background:var(--paper); color:var(--ink); font-family:Inter,ui-sans-serif,system-ui,-apple-system,"Segoe UI",sans-serif; }
.q-layout,.q-page { background:var(--paper); }
.shell { width:min(1800px,calc(100vw - 12px)); margin:0 auto; padding:3px 0 22px; gap:3px; }
.topbar { width:100%; min-height:32px; display:flex; align-items:center; justify-content:space-between; gap:6px; }
.mark { width:28px; height:28px; border-radius:6px; display:grid; place-items:center; background:var(--green); color:#fff; font-size:11px; font-weight:850; }
.compact-intro { width:100%; min-height:18px; display:flex; align-items:center; padding:0 2px; color:#475467; font-size:10px; line-height:1.15; }
.muted { color:var(--muted); }
.eyebrow { color:var(--green); text-transform:uppercase; letter-spacing:.12em; font-size:10px; font-weight:800; }
.pill { border:1px solid var(--line); background:#fff; border-radius:999px; padding:3px 7px; color:var(--muted); font-size:9px; line-height:1.1; }
.workspace { width:100%; display:grid; grid-template-columns:265px minmax(0,1fr); gap:10px; align-items:start; }
.pane { border:1px solid var(--line); background:var(--panel); border-radius:10px; min-width:0; overflow:hidden; }
.pane-head { min-height:38px; padding:6px 9px; border-bottom:1px solid var(--line); background:#fafbfb; }
.template-scroll { max-height:735px; overflow:auto; padding:5px; }
.template-group { width:100%; gap:0; margin-top:3px; }
.template-category { width:100%; min-height:20px; display:flex; align-items:center; padding:2px 5px; margin:0; background:#eef1f3; border:1px solid #cbd2d7; color:#344054; font-size:10px; font-weight:800; line-height:1; letter-spacing:.07em; text-transform:uppercase; }
.template-stack { width:100%; gap:0; }
.template-cell { box-sizing:border-box; width:100%; min-height:42px; height:42px; display:flex; flex-direction:column; align-items:stretch; justify-content:center; gap:0; border:1px solid #b9cfc3; border-top:0; border-left:4px solid var(--green); background:#f7fbf9; padding:3px 7px; cursor:grab; overflow:hidden; }
.template-cell:last-child { border-radius:0 0 6px 6px; }
.template-cell:hover { position:relative; z-index:1; background:#e8f5ed; border-color:#6f9f83; box-shadow:inset 0 0 0 1px #6f9f83; }
.template-title { width:100%; margin:0; color:#173b27; font:700 12px/1.12 Inter,ui-sans-serif,system-ui; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }
.template-code { width:100%; margin:0; color:#1f2937; font:11px/1.12 ui-monospace,SFMono-Regular,Menlo,monospace; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }
.template-copy { display:none; }
.sheet-toolbar { padding:8px; border-bottom:1px solid var(--line); background:#fafbfb; }
.formula-bar { display:grid; grid-template-columns:64px minmax(0,1fr) auto; gap:6px; align-items:center; padding:6px 8px; border-bottom:1px solid var(--line); background:#fff; }
.coord-box { border:1px solid var(--line); background:#f8fafb; min-height:40px; display:grid; place-items:center; font:600 12px ui-monospace,SFMono-Regular,Menlo,monospace; }
.sheet-viewport { width:100%; max-height:735px; overflow:auto; background:#fff; }
.sheet-grid { display:grid; width:max-content; min-width:100%; align-items:stretch; }
.grid-head,.row-head,.sheet-cell { box-sizing:border-box; border-right:1px solid var(--grid); border-bottom:1px solid var(--grid); }
.grid-head { position:sticky; top:0; z-index:4; height:34px; display:grid; place-items:center; background:var(--head); font-size:11px; font-weight:700; color:#52606d; }
.corner { left:0; z-index:6; }
.row-head { position:sticky; left:0; z-index:3; height:78px; background:var(--head); display:flex; flex-direction:column; align-items:center; justify-content:center; font-size:10px; }
.sheet-cell { height:78px; padding:7px 8px; background:#fff; overflow:hidden; cursor:cell; white-space:pre-wrap; overflow-wrap:anywhere; font:11px/1.35 ui-monospace,SFMono-Regular,Menlo,monospace; }
.sheet-cell:hover { outline:2px solid #66a47f; outline-offset:-2px; z-index:2; }
.sheet-cell.target { background:#f7faf8; font-weight:700; color:#254d36; }
.sheet-cell.role { background:#fff7ed; border-top:3px solid #d97706; }
.sheet-cell.reuse { background:#fef3c7; border-top:3px solid #b45309; }
.sheet-cell.brace { background:#f3e8ff; border-top:3px solid #7c3aed; }
.sheet-cell.group { background:#eef2ff; border-top:3px solid #536171; }
.sheet-cell.generate { background:#ecfdf5; border-top:3px solid #0f766e; }
.sheet-cell.empty { color:#a3abb4; }
.cell-first-line { font-weight:650; }
.cell-meta { font:10px/1.3 Inter,ui-sans-serif,system-ui; color:#7a8491; margin-top:4px; }
.sheet-status { padding:7px 10px; border-top:1px solid var(--line); background:#fafbfb; }
.issue-error { border-left:4px solid #dc2626; background:#fef2f2; color:#7f1d1d; padding:7px 9px; }
.issue-warning { border-left:4px solid #d97706; background:#fff7ed; color:#7c2d12; padding:7px 9px; }
.reference-note { border:1px solid #bfd4c7; background:#eef8f2; color:#315c42; border-radius:8px; padding:9px 11px; }
.readonly-wrap { border:1px solid var(--line); background:#fff; border-radius:9px; overflow:auto; max-height:760px; }
.readonly-cell { height:64px; cursor:default; }
.mini-sheets { display:grid; grid-template-columns:repeat(auto-fill,minmax(260px,1fr)); gap:10px; }
.mini-sheet { border:1px solid var(--line); background:#fff; border-radius:8px; overflow:hidden; }
.mini-head { background:var(--head); border-bottom:1px solid var(--line); padding:8px 10px; font-weight:700; }
.mini-row { display:grid; grid-template-columns:52px minmax(0,1fr); border-bottom:1px solid #e5e7eb; min-height:34px; }
.mini-coordinate { background:#f6f7f8; border-right:1px solid #e1e4e7; padding:7px; font:10px ui-monospace,monospace; color:#667085; }
.mini-value { padding:7px; font:11px ui-monospace,monospace; overflow-wrap:anywhere; }
.control-sheet { border:1px solid var(--line); background:#fff; border-radius:9px; overflow:hidden; }
.control-toolbar { min-height:36px; padding:4px 7px; border-bottom:1px solid var(--line); background:#fafbfb; }
.control-grid { display:grid; width:100%; min-width:720px; align-items:stretch; }
.control-grid .row-head,.control-grid .sheet-cell { height:54px; }
.control-grid .sheet-cell { cursor:pointer; padding:6px 8px; }
.control-grid .sheet-cell:hover { outline:2px solid #66a47f; outline-offset:-2px; }
.control-code-grid { grid-template-columns:52px minmax(780px,1fr); }
.control-code-grid .row-head,.control-code-grid .sheet-cell { height:230px; }
.control-code-cell { background:#f8faf9; }
.control-placeholder { color:#98a2ad; font-style:italic; }
.q-tabs { min-height:28px; border-bottom:1px solid var(--line); background:#fff; }
.q-tab { min-height:28px; padding:0 9px; }
.q-tab__label { font-size:10px; line-height:1; letter-spacing:.01em; }
.q-tab--active { color:var(--green)!important; }
.q-tab__indicator { background:var(--green)!important; height:2px; }
.upload-zone { border:1.5px dashed #8eaa99; background:#f4faf6; border-radius:9px; }
.code { white-space:pre-wrap; overflow-wrap:anywhere; background:#172033; color:#e7edf7; border-radius:8px; padding:11px; font:11px/1.5 ui-monospace,SFMono-Regular,Menlo,monospace; }
@media (max-width:920px) { .workspace { grid-template-columns:1fr; } .template-scroll { max-height:330px; } .shell { width:calc(100% - 12px); } .topbar { align-items:flex-start; flex-direction:column; } }
"""


def _heading(title: str, copy: str = "") -> None:
    ui.label(title).classes("font-bold text-sm")
    if copy:
        ui.label(copy).classes("text-xs muted leading-relaxed")


def _drop_value(value: Any) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, list) and value:
        return _drop_value(value[0])
    if isinstance(value, dict):
        for key in ("value", "data", "args"):
            if key in value:
                return _drop_value(value[key])
    return ""


def _cell_kind(value: str, target: bool = False) -> str:
    stripped = value.strip()
    if target:
        return "target" if stripped else "target empty"
    if not stripped:
        return "empty"
    if stripped.startswith("FW_("):
        return "brace"
    if "FW_Reuse" in stripped:
        return "reuse"
    if stripped in {"FW_Exclude", "FW_Optional"}:
        return "role"
    if "FW_Group" in stripped or "FW_ReplaceRE" in stripped:
        return "group"
    return "generate"


def _cell_display(value: str) -> tuple[str, str]:
    if not value.strip():
        return "", "double-click or drop a verb"
    lines = value.splitlines()
    return lines[0], f"+ {len(lines) - 1} lines in this cell" if len(lines) > 1 else ""


def _reference_matrix(analysis: WorkbookAnalysis, sheet_name: str) -> list[list[str]]:
    profile = analysis.sheet_by_name[sheet_name]
    matrix = [["" for _ in range(profile.column_span)] for _ in range(profile.row_span)]
    for cell in profile.values:
        row, column = coordinate_to_tuple(cell.coordinate)
        matrix[row - 1][column - 1] = cell.value
    return matrix


def _readonly_grid(matrix: list[list[str]]) -> None:
    columns = max((len(row) for row in matrix), default=1)
    style = f"grid-template-columns:52px repeat({columns}, 184px)"
    with ui.element("div").classes("readonly-wrap w-full"):
        with ui.element("div").classes("sheet-grid").style(style):
            ui.label("").classes("grid-head corner")
            for column in range(1, columns + 1):
                ui.label(get_column_letter(column)).classes("grid-head")
            for row_index, row in enumerate(matrix, 1):
                ui.label(str(row_index)).classes("row-head")
                for column in range(columns):
                    value = row[column] if column < len(row) else ""
                    first, meta = _cell_display(value)
                    with ui.element("div").classes(f"sheet-cell readonly-cell {_cell_kind(value, target=column == 0)}"):
                        ui.label(first).classes("cell-first-line")
                        if meta:
                            ui.label(meta).classes("cell-meta")


def _sheet_names_view(project: GridProject) -> None:
    rows = project.sequence_rows
    matrix = [[row.target, row.prefix or "FW_EMPTY_STRING", row.suffix or "FW_EMPTY_STRING"] for row in rows]
    with ui.column().classes("w-full gap-3"):
        ui.label("Derived live from the composition grid. Edit target, prefix, and suffix by double-clicking column A in FW_Seq.").classes("reference-note text-xs w-full")
        _readonly_grid(matrix)


def _data_sheets_view(project: GridProject) -> None:
    with ui.column().classes("w-full gap-3"):
        ui.label("Source sheets are deliberately separate from directive cells. Double-click a target cell in FW_Seq to edit one of these sheets.").classes("reference-note text-xs w-full")
        with ui.element("div").classes("mini-sheets w-full"):
            for row in project.sequence_rows:
                with ui.element("div").classes("mini-sheet"):
                    with ui.element("div").classes("mini-head"):
                        ui.label(row.target or "Unnamed target")
                    for index, value in enumerate(row.values, 1):
                        with ui.element("div").classes("mini-row"):
                            ui.label(f"A{index}").classes("mini-coordinate")
                            ui.label(value).classes("mini-value")
            for name, values in project.auxiliary_sheets.items():
                with ui.element("div").classes("mini-sheet"):
                    with ui.element("div").classes("mini-head"):
                        ui.label(f"{name} · helper")
                    for index, value in enumerate(values, 1):
                        with ui.element("div").classes("mini-row"):
                            ui.label(f"A{index}").classes("mini-coordinate")
                            ui.label(value).classes("mini-value")


def _editable_control_cell(value: str, on_open: Any, placeholder: str = "Double-click to edit") -> None:
    first, meta = _cell_display(value) if value else ("", placeholder)
    cell = ui.element("div").classes("sheet-cell control-cell")
    cell.on("dblclick", on_open)
    with cell:
        ui.tooltip(placeholder)
        ui.label(first).classes("cell-first-line")
        if meta:
            ui.label(meta).classes("cell-meta")


def _project_info_matrix(project: GridProject) -> list[list[str]]:
    return [
        [row.target] + [row.directive(index) for index in range(project.directive_columns)]
        for row in project.sequence_rows
    ]


def _project_info_view(project: GridProject) -> None:
    with ui.column().classes("w-full gap-1"):
        ui.label("Live FW_Info mirror of the composed FW_Seq. This is the FW_Info sheet written to the downloaded workbook.").classes("reference-note text-xs w-full")
        _readonly_grid(_project_info_matrix(project))


def _run_once_view(project: GridProject, on_project_changed: Any) -> None:
    def open_editor() -> None:
        with ui.dialog() as dialog, ui.card().classes("w-[1000px] max-w-[97vw] max-h-[95vh] p-4 gap-3"):
            _heading("FW_RunMeFirstOnce · A1", "Core reads every populated cell as bootstrap code; Face 1 new uses the canonical single A1 Java source cell.")
            editor = ui.textarea(value=project.run_once_source, label="Exact A1 source").props("outlined autogrow input-class=font-mono").classes("w-full max-h-[68vh] overflow-auto")
            ui.label("Expected by Core: class RunMeFirstOnce, public static String FW_ARGS, and an FW_ARGS assignment.").classes("text-[10px] muted")

            def clear() -> None:
                project.run_once_source = ""
                ui.notify("Cleared FW_RunMeFirstOnce A1")
                dialog.close()
                on_project_changed()

            def save() -> None:
                project.run_once_source = str(editor.value or "")
                ui.notify(f"Saved FW_RunMeFirstOnce A1 · {len(project.run_once_source)} characters", type="positive")
                dialog.close()
                on_project_changed()

            with ui.row().classes("w-full justify-between gap-2"):
                ui.button("Clear A1", icon="backspace", on_click=clear).props("flat color=red-7")
                with ui.row().classes("gap-2"):
                    ui.button("Cancel", on_click=dialog.close).props("flat")
                    ui.button("Save A1", icon="save", on_click=save).props("unelevated color=green-8")
        dialog.open()

    with ui.column().classes("control-sheet w-full gap-0"):
        with ui.row().classes("control-toolbar w-full items-center justify-between"):
            ui.label("One physical source cell · double-click A1 to expand").classes("text-[10px] muted")
            ui.button("Edit A1", icon="code", on_click=open_editor).props("flat dense no-caps color=green-8")
        with ui.element("div").classes("control-grid control-code-grid"):
            ui.label("").classes("grid-head corner")
            ui.label("A · Java bootstrap source").classes("grid-head")
            ui.label("1").classes("row-head")
            _editable_control_cell(project.run_once_source, open_editor, "Double-click A1 to add bootstrap source")


def _arguments_view(project: GridProject, on_project_changed: Any) -> None:
    def open_editor(index: int) -> None:
        exists = index < len(project.arguments)
        current = project.arguments[index] if exists else ""
        with ui.dialog() as dialog, ui.card().classes("w-[680px] max-w-[96vw] p-4 gap-3"):
            _heading(f"FW_Arguments · A{index + 1}", "Core inserts every populated cell as one argument value. Empty text exports as FW_EMPTY_STRING.")
            editor = ui.input("Exact argument value", value=current).props("outlined").classes("w-full font-mono")

            def delete() -> None:
                if exists:
                    project.arguments.pop(index)
                ui.notify(f"Deleted FW_Arguments row {index + 1}")
                dialog.close()
                on_project_changed()

            def save() -> None:
                value = str(editor.value or "")
                if exists:
                    project.arguments[index] = value
                else:
                    project.arguments.append(value)
                ui.notify(f"Saved FW_Arguments A{index + 1}", type="positive")
                dialog.close()
                on_project_changed()

            with ui.row().classes("w-full justify-between gap-2"):
                if exists:
                    ui.button("Delete row", icon="delete_outline", on_click=delete).props("flat color=red-7")
                else:
                    ui.space()
                with ui.row().classes("gap-2"):
                    ui.button("Cancel", on_click=dialog.close).props("flat")
                    ui.button("Save argument", icon="save", on_click=save).props("unelevated color=green-8")
        dialog.open()

    with ui.column().classes("control-sheet w-full gap-0"):
        with ui.row().classes("control-toolbar w-full items-center justify-between"):
            ui.label(f"{len(project.arguments)} argument cell(s) · one value per row").classes("text-[10px] muted")
            ui.button("Add argument", icon="add", on_click=lambda: open_editor(len(project.arguments))).props("flat dense no-caps color=green-8")
        with ui.element("div").classes("control-grid").style("grid-template-columns:52px minmax(720px,1fr)"):
            ui.label("").classes("grid-head corner")
            ui.label("A · argument").classes("grid-head")
            values = project.arguments if project.arguments else [""]
            for index, value in enumerate(values):
                ui.label(str(index + 1)).classes("row-head")
                shown = value if value else ("FW_EMPTY_STRING" if project.arguments else "")
                _editable_control_cell(shown, partial(open_editor, index), "Double-click to add or edit this argument")


def _custom_var_view(project: GridProject, on_project_changed: Any) -> None:
    def open_editor(index: int) -> None:
        exists = index < len(project.custom_verdicts)
        current_code, current_message = project.custom_verdicts[index] if exists else ("", "")
        with ui.dialog() as dialog, ui.card().classes("w-[760px] max-w-[96vw] p-4 gap-3"):
            _heading(f"FW_CUSTOM_VAR · row {index + 1}", "Column A is a Java Integer verdict code; column B is its message.")
            code = ui.input("A · verdict code", value=current_code).props("outlined").classes("w-full font-mono")
            message = ui.textarea("B · verdict message", value=current_message).props("outlined autogrow").classes("w-full")

            def delete() -> None:
                if exists:
                    project.custom_verdicts.pop(index)
                ui.notify(f"Deleted FW_CUSTOM_VAR row {index + 1}")
                dialog.close()
                on_project_changed()

            def save() -> None:
                code_text = str(code.value or "").strip()
                try:
                    parsed = int(code_text)
                except ValueError:
                    ui.notify("Verdict code must be a Java integer", type="negative")
                    return
                if not -(2**31) <= parsed < 2**31:
                    ui.notify("Verdict code is outside Java Integer range", type="negative")
                    return
                for row, (item_code, _) in enumerate(project.custom_verdicts):
                    if row == index:
                        continue
                    try:
                        existing_code = int(str(item_code).strip())
                    except ValueError:
                        continue
                    if existing_code == parsed:
                        ui.notify(f"Verdict code {parsed} already exists", type="negative")
                        return
                value = (str(parsed), str(message.value or ""))
                if exists:
                    project.custom_verdicts[index] = value
                else:
                    project.custom_verdicts.append(value)
                ui.notify(f"Saved FW_CUSTOM_VAR row {index + 1}", type="positive")
                dialog.close()
                on_project_changed()

            with ui.row().classes("w-full justify-between gap-2"):
                if exists:
                    ui.button("Delete row", icon="delete_outline", on_click=delete).props("flat color=red-7")
                else:
                    ui.space()
                with ui.row().classes("gap-2"):
                    ui.button("Cancel", on_click=dialog.close).props("flat")
                    ui.button("Save verdict", icon="save", on_click=save).props("unelevated color=green-8")
        dialog.open()

    with ui.column().classes("control-sheet w-full gap-0"):
        with ui.row().classes("control-toolbar w-full items-center justify-between"):
            ui.label(f"{len(project.custom_verdicts)} custom verdict row(s) · double-click either cell to edit the row").classes("text-[10px] muted")
            ui.button("Add verdict", icon="add", on_click=lambda: open_editor(len(project.custom_verdicts))).props("flat dense no-caps color=green-8")
        with ui.element("div").classes("control-grid").style("grid-template-columns:52px 180px minmax(540px,1fr)"):
            ui.label("").classes("grid-head corner")
            ui.label("A · Integer code").classes("grid-head")
            ui.label("B · message").classes("grid-head")
            values = project.custom_verdicts if project.custom_verdicts else [("", "")]
            for index, (code, message) in enumerate(values):
                ui.label(str(index + 1)).classes("row-head")
                _editable_control_cell(code, partial(open_editor, index), "Double-click to add or edit verdict")
                shown_message = message if message else ("FW_EMPTY_STRING" if project.custom_verdicts else "")
                _editable_control_cell(shown_message, partial(open_editor, index), "Double-click to add or edit verdict")


def _fw_info_example_view(analysis: WorkbookAnalysis) -> None:
    with ui.column().classes("w-full gap-3"):
        ui.label("Read-only FW_Info from example150726.xlsx. It is a positional grammar example, not content copied into a new project.").classes("reference-note text-xs w-full")
        _readonly_grid(_reference_matrix(analysis, "FW_Info"))


def _sequence_editor(state: dict[str, Any], on_project_changed: Any) -> None:
    project: GridProject = state["project"]
    selection: dict[str, Any] = {"row_id": project.rows[0].id, "column": 0}
    grid_host: Any = None

    def row_number(row_id: str) -> int:
        return next((index + 1 for index, row in enumerate(project.rows) if row.id == row_id), 1)

    def coordinate(row_id: str, column: int) -> str:
        return f"{get_column_letter(column + 1)}{row_number(row_id)}"

    def select_cell(row_id: str, column: int) -> None:
        selection["row_id"] = row_id
        selection["column"] = column
        coordinate_label.set_text(coordinate(row_id, column))
        formula.set_value(project.cell(row_id, column))
        formula_hint.set_text("target sheet + data" if column == 0 else "exact FW_Seq cell; newlines stay inside this XLSX cell")

    def refresh_grid() -> None:
        grid_host.clear()
        render_grid()
        row_id = selection.get("row_id")
        if project.row(row_id) is None:
            selection.update(row_id=project.rows[0].id, column=0)
        select_cell(selection["row_id"], selection["column"])

    def sync_views() -> None:
        refresh_grid()
        on_project_changed()

    def save_formula() -> None:
        project.set_cell(selection["row_id"], selection["column"], str(formula.value or ""))
        sync_views()
        ui.notify(f"Saved {coordinate(selection['row_id'], selection['column'])}", type="positive")

    def clear_directive(row_id: str, column: int, dialog) -> None:
        project.set_cell(row_id, column, "")
        dialog.close()
        selection.update(row_id=row_id, column=column)
        sync_views()

    def open_directive_editor(row_id: str, column: int) -> None:
        value = project.cell(row_id, column)
        with ui.dialog() as dialog, ui.card().classes("w-[760px] max-w-[96vw] p-5 gap-4"):
            with ui.row().classes("w-full items-start justify-between gap-3"):
                with ui.column().classes("gap-0"):
                    _heading(f"Edit {coordinate(row_id, column)}", "One grid cell equals one physical FW_Seq XLSX cell. Multi-line Group/ReplaceRE pipelines remain in this cell.")
                ui.label("directive cell").classes("pill")
            editor = ui.textarea(value=value, label="Exact cell content").props("outlined autogrow input-class=font-mono").classes("w-full")
            ui.label("Column position is significant and will not be normalized. FW_Group may be in the middle of a row.").classes("text-xs muted")
            with ui.row().classes("w-full justify-between gap-2"):
                ui.button("Clear cell", icon="backspace", on_click=partial(clear_directive, row_id, column, dialog)).props("flat color=red-7")
                with ui.row().classes("gap-2"):
                    ui.button("Cancel", on_click=dialog.close).props("flat")

                    def save() -> None:
                        project.set_cell(row_id, column, str(editor.value or "").strip())
                        dialog.close()
                        selection.update(row_id=row_id, column=column)
                        sync_views()

                    ui.button("Save cell", icon="save", on_click=save).props("unelevated color=green-8")
        dialog.open()

    def open_target_editor(row_id: str) -> None:
        row = project.row(row_id)
        if row is None:
            return
        draft = list(row.values)
        with ui.dialog() as dialog, ui.card().classes("w-[780px] max-w-[96vw] max-h-[94vh] p-5 gap-4"):
            with ui.row().classes("w-full items-start justify-between gap-3"):
                _heading(f"Target and data sheet · row {row_number(row_id)}", "Column A names the source sheet; its cell data lives in A1, A2, A3… of that sheet.")
                ui.label(coordinate(row_id, 0)).classes("pill font-mono")
            target = ui.input("Target sheet name", value=row.target).props("outlined").classes("w-full")
            with ui.row().classes("w-full gap-2"):
                prefix = ui.input("Output prefix", value=row.prefix).props("outlined dense").classes("grow")
                suffix = ui.input("Output suffix", value=row.suffix).props("outlined dense").classes("grow")
            ui.label("Data sheet cells").classes("eyebrow")
            rows_host = ui.column().classes("w-full gap-2 overflow-auto max-h-[48vh]")

            def set_value(index: int, e: events.ValueChangeEventArguments) -> None:
                draft[index] = str(e.value or "")

            def remove_value(index: int) -> None:
                draft.pop(index)
                render_values()

            def render_values() -> None:
                rows_host.clear()
                with rows_host:
                    if not draft:
                        ui.label("No data cells. Add A1 to continue.").classes("text-xs muted py-4 text-center w-full")
                    for index, value in enumerate(draft):
                        with ui.row().classes("w-full items-center gap-2 no-wrap"):
                            ui.label(f"A{index + 1}").classes("coord-box w-14")
                            cell = ui.input(value=value).props("outlined dense").classes("grow font-mono")
                            cell.on_value_change(partial(set_value, index))
                            ui.button(icon="delete_outline", on_click=partial(remove_value, index)).props("flat round dense color=red-7")

            def add_value() -> None:
                draft.append("")
                render_values()

            def delete_grid_row() -> None:
                project.rows = [item for item in project.rows if item.id != row_id]
                if not project.rows:
                    project.add_blank_rows(1)
                dialog.close()
                selection.update(row_id=project.rows[0].id, column=0)
                sync_views()

            def save() -> None:
                row.target = str(target.value or "").strip()
                row.prefix = str(prefix.value or "")
                row.suffix = str(suffix.value or "")
                row.values = list(draft)
                ui.notify(f"Saved target and {len(draft)} data cells", type="positive")
                dialog.close()
                selection.update(row_id=row_id, column=0)
                sync_views()

            render_values()
            with ui.row().classes("w-full items-center justify-between gap-2 flex-wrap"):
                with ui.row().classes("gap-2"):
                    ui.button("Add data row", icon="add", on_click=add_value).props("flat color=green-8")
                    ui.button("Delete FW_Seq row", icon="delete", on_click=delete_grid_row).props("flat color=red-7")
                with ui.row().classes("gap-2"):
                    ui.button("Cancel", on_click=dialog.close).props("flat")
                    ui.button("Save target + data", icon="save", on_click=save).props("unelevated color=green-8")
        dialog.open()

    def open_cell_editor(row_id: str, column: int) -> None:
        select_cell(row_id, column)
        if column == 0:
            open_target_editor(row_id)
        else:
            open_directive_editor(row_id, column)

    def drop_template(row_id: str, column: int, e: events.GenericEventArguments) -> None:
        template_id = _drop_value(e.args)
        template = TEMPLATE_BY_ID.get(template_id)
        if template is None:
            return
        if column == 0:
            ui.notify("Verb templates belong in directive columns B onward", type="warning")
            return
        if project.cell(row_id, column).strip():
            select_cell(row_id, column)
            ui.notify(f"{coordinate(row_id, column)} is occupied; double-click it to edit or clear", type="warning")
            return
        directive = template.directive
        if directive.startswith("FW_(") and directive.endswith(")"):
            row_index = next((index for index, item in enumerate(project.rows) if item.id == row_id), -1)
            fields = directive[4:-1].split(",")
            if row_index >= 2 and len(fields) == 9:
                if fields[2] == "LEFT":
                    fields[2] = project.rows[row_index - 2].target.strip()
                if fields[4] == "RIGHT":
                    fields[4] = project.rows[row_index - 1].target.strip()
                directive = f"FW_({','.join(fields)})"
        project.set_cell(row_id, column, directive)
        selection.update(row_id=row_id, column=column)
        sync_views()

    def move_row(row_id: str, delta: int) -> None:
        project.move_row(row_id, delta)
        selection.update(row_id=row_id, column=selection.get("column", 0))
        sync_views()

    def validate_project() -> None:
        refresh_grid()
        issues = project.validate()
        errors = sum(issue.severity == "error" for issue in issues)
        warnings = sum(issue.severity == "warning" for issue in issues)
        if errors:
            ui.notify(f"{errors} export-blocking error(s), {warnings} warning(s)", type="negative")
        elif warnings:
            ui.notify(f"Valid with {warnings} warning(s)", type="warning")
        else:
            ui.notify("FW_Seq positional contract is valid", type="positive")

    def export_project() -> None:
        try:
            content = project.to_xlsx_bytes()
        except ValueError as exc:
            ui.notify(str(exc), type="negative", multi_line=True)
            refresh_grid()
            return
        filename = (project.name.strip() or "fw_sequence") + ".xlsx"
        ui.download(content, filename=filename, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def add_helper() -> None:
        with ui.dialog() as dialog, ui.card().classes("w-[620px] max-w-[95vw] p-5 gap-4"):
            _heading("Add helper data sheet", "Reference this name from FW_Cartes, FW_Cartes_first, or FW_Separator.")
            name = ui.input("Sheet name", value="SEPARATOR").props("outlined").classes("w-full")
            values = ui.textarea("Column A values · one per line", value=", ").props("outlined autogrow").classes("w-full")
            with ui.row().classes("w-full justify-end gap-2"):
                ui.button("Cancel", on_click=dialog.close).props("flat")

                def save() -> None:
                    sheet_name = str(name.value or "").strip()
                    if not sheet_name:
                        ui.notify("Helper sheet needs a name", type="negative")
                        return
                    project.auxiliary_sheets[sheet_name] = str(values.value or "").splitlines()
                    dialog.close()
                    sync_views()

                ui.button("Save helper", icon="save", on_click=save).props("unelevated color=green-8")
        dialog.open()

    def render_grid() -> None:
        issues = project.validate()
        issue_cells = {(issue.row_id, issue.column) for issue in issues if issue.row_id}
        style = f"grid-template-columns:52px 190px repeat({project.directive_columns}, 190px)"
        with grid_host:
            with ui.element("div").classes("sheet-viewport"):
                with ui.element("div").classes("sheet-grid").style(style):
                    ui.label("").classes("grid-head corner")
                    for column in range(project.directive_columns + 1):
                        label = get_column_letter(column + 1)
                        if column == 0:
                            label += " · target/data"
                        ui.label(label).classes("grid-head")
                    for index, row in enumerate(project.rows, 1):
                        with ui.element("div").classes("row-head"):
                            ui.label(str(index)).classes("font-bold")
                            with ui.row().classes("gap-0 no-wrap"):
                                ui.button(icon="keyboard_arrow_up", on_click=partial(move_row, row.id, -1)).props("flat dense round size=xs")
                                ui.button(icon="keyboard_arrow_down", on_click=partial(move_row, row.id, 1)).props("flat dense round size=xs")
                        for column in range(project.directive_columns + 1):
                            value = project.cell(row.id, column)
                            first, meta = _cell_display(value)
                            if column == 0 and value:
                                meta = f"{len(row.values)} data cells · double-click to expand"
                            classes = f"sheet-cell {_cell_kind(value, target=column == 0)}"
                            if (row.id, column) in issue_cells or (row.id, None) in issue_cells:
                                classes += " !border-t-red-600"
                            cell = ui.element("div").classes(classes)
                            cell.on("click", partial(select_cell, row.id, column))
                            cell.on("dblclick", partial(open_cell_editor, row.id, column))
                            if column > 0:
                                cell.on("dragover", js_handler="(event) => { event.preventDefault(); event.dataTransfer.dropEffect = 'copy'; }")
                                cell.on("drop", partial(drop_template, row.id, column), js_handler="(event) => { event.preventDefault(); emit(event.dataTransfer.getData('text/plain')); }")
                            with cell:
                                ui.label(first).classes("cell-first-line")
                                if meta:
                                    ui.label(meta).classes("cell-meta")
            with ui.row().classes("sheet-status w-full items-center justify-between gap-3 flex-wrap"):
                ui.label(f"{len(project.sequence_rows)} used rows · {project.directive_columns + 1} columns · trailing blank rows are not exported").classes("text-[11px] muted")
                ui.label(f"{sum(i.severity == 'error' for i in issues)} errors · {sum(i.severity == 'warning' for i in issues)} warnings").classes("text-[11px] font-semibold")
            if issues:
                with ui.expansion("Validation details", icon="fact_check").classes("w-full border-t"):
                    for issue in issues[:30]:
                        ui.label(issue.message).classes(f"issue-{issue.severity} text-xs w-full mb-1")
                    if len(issues) > 30:
                        ui.label(f"+ {len(issues) - 30} more issues").classes("text-xs muted")

    with ui.element("div").classes("workspace"):
        with ui.column().classes("pane gap-0"):
            with ui.element("div").classes("pane-head"):
                _heading("Verb templates", "Drag one cell into an exact B… column")
            with ui.column().classes("template-scroll w-full gap-0"):
                ui.button("Insert brace trio", icon="account_tree", on_click=lambda: (project.append_brace_trio(), sync_views())).props("outline color=deep-purple-7 dense no-caps").classes("w-full text-[11px]")
                for category in dict.fromkeys(template.category for template in TEMPLATES):
                    with ui.column().classes("template-group"):
                        ui.label(category).classes("template-category")
                        with ui.column().classes("template-stack"):
                            for template in (item for item in TEMPLATES if item.category == category):
                                tile = ui.element("div").classes("template-cell")
                                tile.props("draggable=true")
                                tile.on("dragstart", js_handler=f"(event) => {{ event.dataTransfer.setData('text/plain', '{template.id}'); event.dataTransfer.effectAllowed = 'copy'; }}")
                                with tile:
                                    ui.tooltip(f"{template.directive}\n{template.description}")
                                    ui.label(template.label).classes("template-title")
                                    ui.label(template.directive.splitlines()[0]).classes("template-code")

        with ui.column().classes("pane gap-0"):
            with ui.row().classes("sheet-toolbar w-full items-center justify-between gap-2 flex-wrap"):
                name = ui.input("Project", value=project.name).props("outlined dense").classes("w-56")

                def update_name(e: events.ValueChangeEventArguments) -> None:
                    project.name = str(e.value or "")
                    on_project_changed()

                name.on_value_change(update_name)
                with ui.row().classes("gap-1 flex-wrap"):
                    ui.button("+5 rows", icon="add", on_click=lambda: (project.add_blank_rows(5), sync_views())).props("flat color=green-8")
                    ui.button("+ column", icon="view_column", on_click=lambda: (project.add_directive_column(), sync_views())).props("flat color=green-8")
                    ui.button("Helper", icon="post_add", on_click=add_helper).props("flat color=green-8")
                    ui.button("Validate", icon="fact_check", on_click=validate_project).props("outline color=green-8")
                    ui.button("Download XLSX", icon="download", on_click=export_project).props("unelevated color=green-8")
            with ui.element("div").classes("formula-bar w-full"):
                coordinate_label = ui.label("A1").classes("coord-box")
                formula = ui.input(value=project.cell(selection["row_id"], 0)).props("outlined dense").classes("w-full font-mono")
                ui.button("Save cell", icon="save", on_click=save_formula).props("flat color=green-8")
            formula_hint = ui.label("target sheet + data").classes("text-[10px] muted px-2 py-1 border-b w-full")
            grid_host = ui.column().classes("w-full gap-0")
            render_grid()


@ui.page("/")
def index() -> None:
    ui.add_css(CSS + RUN_CSS + WIZARD_CSS)
    ui.page_title("Face 1 new · FW_Seq sheet composer")
    reference = analyze_workbook(EXAMPLE_WORKBOOK)
    initial_project = GridProject.blank()
    state: dict[str, Any] = {
        "project": initial_project,
        "reference": reference,
        "runtime_config": default_run_config(initial_project.name),
    }
    workspace = ui.column().classes("shell")

    def poll_runtime() -> None:
        callback = state.get("poll_runtime")
        if callback is not None:
            callback()

    ui.timer(1.0, poll_runtime)

    with ui.dialog() as upload_dialog, ui.card().classes("w-[630px] max-w-[95vw] p-5 gap-4"):
        _heading("Import a workbook into the grid", "Import is explicit. The bundled example is otherwise reference-only and is not copied into a new composition.")
        upload_status = ui.label("Choose an .xlsx file up to 12 MB").classes("text-xs muted")
        upload = ui.upload(label="Drop XLSX here or browse", auto_upload=True, max_file_size=12_000_000).props("accept=.xlsx flat bordered").classes("upload-zone w-full")
        with ui.row().classes("w-full justify-end"):
            ui.button("Close", on_click=upload_dialog.close).props("flat")

    def new_project() -> None:
        session = state.get("session")
        if session is not None and not session.done:
            ui.notify("Cancel or finish the active run before replacing the workbook", type="warning")
            return
        state["project"] = GridProject.blank()
        state["runtime_config"] = default_run_config(state["project"].name)
        ui.notify("New minimal six-sheet Core workbook created")
        rebuild()

    def apply_guided_project(project: GridProject) -> None:
        session = state.get("session")
        if session is not None and not session.done:
            ui.notify("Cancel or finish the active run before applying a guided starter", type="warning")
            return
        state["project"] = project
        state["runtime_config"] = default_run_config(project.name)
        rebuild()
        ui.notify("Guided starter applied. Review and edit every proposed workbook cell.", type="positive")

    def rebuild() -> None:
        project: GridProject = state["project"]
        state.pop("poll_runtime", None)
        state.pop("refresh_runtime", None)
        workspace.clear()
        view_hosts: dict[str, Any] = {}

        def refresh_project_views() -> None:
            renderers = {
                "names": lambda: _sheet_names_view(project),
                "run_once": lambda: _run_once_view(project, refresh_project_views),
                "arguments": lambda: _arguments_view(project, refresh_project_views),
                "custom_var": lambda: _custom_var_view(project, refresh_project_views),
                "info": lambda: _project_info_view(project),
                "data": lambda: _data_sheets_view(project),
            }
            for name, renderer in renderers.items():
                host = view_hosts.get(name)
                if host is None:
                    continue
                host.clear()
                with host:
                    renderer()
            refresh_runtime = state.get("refresh_runtime")
            if refresh_runtime is not None:
                refresh_runtime()

        with workspace:
            with ui.element("div").classes("topbar"):
                with ui.row().classes("items-center gap-2 no-wrap"):
                    ui.html('<div class="mark">F1+</div>', sanitize=False)
                    with ui.column().classes("gap-0"):
                        ui.label("Face 1 new").classes("font-bold text-xs leading-none")
                        ui.label("XLSX-shaped Core workbook composition").classes("text-[9px] muted leading-none")
                with ui.row().classes("items-center gap-1 flex-wrap"):
                    ui.label("example150726.xlsx · reference only").classes("pill")
                    ui.button("New", on_click=new_project).props("flat dense no-caps size=sm color=green-8")
                    ui.button("Import XLSX", on_click=upload_dialog.open).props("outline dense no-caps size=sm color=green-8")
            ui.label("Compose FW_Seq and all six Core control sheets; double-click a sheet cell to expand its editor.").classes("compact-intro")

            with ui.tabs().props("dense align=left outside-arrows mobile-arrows").classes("w-full") as tabs:
                invitation_tab = ui.tab("Guided start")
                seq_tab = ui.tab("FW_Seq")
                names_tab = ui.tab("FW_SheetNames")
                run_once_tab = ui.tab("FW_RunMeFirstOnce")
                arguments_tab = ui.tab("FW_Arguments")
                custom_var_tab = ui.tab("FW_CUSTOM_VAR")
                info_tab = ui.tab("FW_Info")
                data_tab = ui.tab("Data sheets")
                run_tab = ui.tab("Plan & run")
                example_info_tab = ui.tab("Example · FW_Info")
            with ui.tab_panels(tabs, value=invitation_tab).classes("w-full bg-transparent p-0"):
                with ui.tab_panel(invitation_tab).classes("px-0 py-0"):
                    render_invitation_center(state, apply_guided_project)
                with ui.tab_panel(seq_tab).classes("px-0 py-0"):
                    _sequence_editor(state, refresh_project_views)
                with ui.tab_panel(names_tab).classes("px-0 py-0"):
                    view_hosts["names"] = ui.column().classes("w-full gap-0")
                with ui.tab_panel(run_once_tab).classes("px-0 py-0"):
                    view_hosts["run_once"] = ui.column().classes("w-full gap-0")
                with ui.tab_panel(arguments_tab).classes("px-0 py-0"):
                    view_hosts["arguments"] = ui.column().classes("w-full gap-0")
                with ui.tab_panel(custom_var_tab).classes("px-0 py-0"):
                    view_hosts["custom_var"] = ui.column().classes("w-full gap-0")
                with ui.tab_panel(info_tab).classes("px-0 py-0"):
                    view_hosts["info"] = ui.column().classes("w-full gap-0")
                with ui.tab_panel(data_tab).classes("px-0 py-0"):
                    view_hosts["data"] = ui.column().classes("w-full gap-0")
                with ui.tab_panel(run_tab).classes("px-0 py-0"):
                    render_run_center(state)
                with ui.tab_panel(example_info_tab).classes("px-0 py-0"):
                    _fw_info_example_view(reference)
            refresh_project_views()
            ui.label("Face 1 remains untouched in intake/. Face 1 new authors the exact Core workbook and can hand that workbook to the real Bundle pipeline explicitly from Plan & run.").classes("text-[10px] muted text-center w-full")

    async def handle_upload(e: events.UploadEventArguments) -> None:
        session = state.get("session")
        if session is not None and not session.done:
            ui.notify("Cancel or finish the active run before importing another workbook", type="warning")
            return
        if not e.file.name.lower().endswith(".xlsx"):
            ui.notify("Please choose an XLSX workbook", type="warning")
            return
        upload_status.set_text(f"Reading {e.file.name}…")
        try:
            raw = await e.file.read()
            analysis = analyze_workbook(raw, source_name=e.file.name)
        except WorkbookInputError as exc:
            upload_status.set_text(str(exc))
            ui.notify(str(exc), type="negative", multi_line=True)
            return
        state["project"] = project_from_analysis(analysis)
        state["runtime_config"] = default_run_config(state["project"].name)
        upload_status.set_text("Workbook placed into the grid.")
        upload_dialog.close()
        rebuild()
        ui.notify(f"Imported {e.file.name}", type="positive")

    upload.on_upload(handle_upload)
    rebuild()


def main() -> None:
    parser = argparse.ArgumentParser(description="Launch the Face 1 new FW_Seq spreadsheet composer")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", default=8088, type=int)
    parser.add_argument("--no-browser", action="store_true")
    parser.add_argument("--reload", action="store_true")
    args = parser.parse_args()
    ui.run(host=args.host, port=args.port, title="Face 1 new · FW_Seq sheet composer", show=not args.no_browser, reload=args.reload, language="en-US")


if __name__ in {"__main__", "__mp_main__"}:
    main()
