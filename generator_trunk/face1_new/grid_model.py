"""Spreadsheet-shaped FW_Seq composition model for Face 1 new."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
import re
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .workbook_model import CONTROL_SHEETS, WorkbookAnalysis


EMPTY = "FW_EMPTY_STRING"
INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


@dataclass(frozen=True)
class CellTemplate:
    id: str
    label: str
    category: str
    directive: str
    description: str
    color: str


GROUP_CLEANUP = (
    'FW_Group\n'
    'FW_ReplaceRE("^\\[", "")\n'
    'FW_ReplaceRE("\\]$", "")\n'
    'FW_ReplaceRE("\\[\\]", "[" + 0 +"]")\n'
    'FW_ReplaceRE("\\]\\, \\[", ", ")\n'
    'FW_ReplaceRE("\\[", "")\n'
    'FW_ReplaceRE("\\]", "")'
)


TEMPLATES: tuple[CellTemplate, ...] = (
    CellTemplate("exclude", "FW_Exclude", "Row role", "FW_Exclude", "Operand/intermediate row. Put before Reuse and generative cells.", "#b45309"),
    CellTemplate("reuse", "FW_Reuse", "Row role", "FW_Reuse", "Only immediately after Exclude on both rows directly before a brace row.", "#b45309"),
    CellTemplate("reuse_table", "FW_ReuseTableOnly", "Row role", "FW_ReuseTableOnly", "Table-only variant; same strict location as Reuse.", "#b45309"),
    CellTemplate("optional", "FW_Optional", "Row role", "FW_Optional", "Reader assembles this result as optional.", "#b45309"),
    CellTemplate("heading", "FW_Heading", "Row role", "FW_Heading", "Excluded row that also acts as a heading.", "#b45309"),
    CellTemplate("last_queue", "FW_LastInQueue", "Row role", "FW_LastInQueue", "Move this source to the end of the queue.", "#b45309"),
    CellTemplate("concatenator", "FW_Concatenator", "Row role", "FW_Concatenator=", "Set the output concatenator after the equals sign.", "#b45309"),
    CellTemplate("combi_bare", "Choose · default", "Choose", "FW_Combi", "Bare form; Core resolves m to 1.", "#0f766e"),
    CellTemplate("combi_default", "Choose · ()", "Choose", "FW_Combi()", "Empty parentheses; Core resolves m to 1.", "#0f766e"),
    CellTemplate("combi_zero", "Choose · zero", "Choose", "FW_Combi(0)", "Numeric zero-size form accepted by Core.", "#0f766e"),
    CellTemplate("combi_one", "Choose · one", "Choose", "FW_Combi(1)", "One source value.", "#0f766e"),
    CellTemplate("combi_k", "Choose · k", "Choose", "FW_Combi(2)", "Unordered k-of-n selection; edit k.", "#0f766e"),
    CellTemplate("combi_size", "Choose · size", "Choose", "FW_Combi(size)", "Use the current source size.", "#0f766e"),
    CellTemplate("combi_all", "Choose · all", "Choose", "FW_Combi(all)", "Run every supported non-empty size.", "#0f766e"),
    CellTemplate("combi_full", "Choose · full", "Choose", "FW_Combi(full)", "Core alias of the all-sizes form.", "#0f766e"),
    CellTemplate("combi_repeat", "ChooseR · k", "Choose", "FW_CombiR(2)", "Combination with repetition.", "#0f766e"),
    CellTemplate("combi_repeat_size", "ChooseR · size", "Choose", "FW_CombiR(size)", "Repeated combination at source size.", "#0f766e"),
    CellTemplate("combi_repeat_all", "ChooseR · all", "Choose", "FW_CombiR(all)", "Repeated combinations for all sizes.", "#0f766e"),
    CellTemplate("combi_repeat_full", "ChooseR · full", "Choose", "FW_CombiR(full)", "Core alias of repeated all-sizes form.", "#0f766e"),
    CellTemplate("permut", "Permute all", "Order", "FW_Permut", "Every ordering of all current atoms.", "#3159c6"),
    CellTemplate("permut_empty", "Permute · ()", "Order", "FW_Permut()", "Empty parentheses; Core resolves m to 1.", "#3159c6"),
    CellTemplate("permut_k", "Permute · k", "Order", "FW_Permut(2)", "Numeric parameter form accepted by Core.", "#3159c6"),
    CellTemplate("permut_identical", "Permute · duplicates", "Order", "FW_Permut(PermutationGenerator.TreatDuplicatesAs.IDENTICAL)", "Treat duplicate inputs as identical.", "#3159c6"),
    CellTemplate("permut_repeat", "PermuteR · k", "Order", "FW_PermutR(2)", "Ordered selection with repetition.", "#3159c6"),
    CellTemplate("permut_repeat_size", "PermuteR · size", "Order", "FW_PermutR(size)", "Repeated permutation at source size.", "#3159c6"),
    CellTemplate("permut_repeat_all", "PermuteR · all", "Order", "FW_PermutR(all)", "Repeated permutations for all sizes.", "#3159c6"),
    CellTemplate("permut_repeat_full", "PermuteR · full", "Order", "FW_PermutR(full)", "Core alias of repeated all-sizes form.", "#3159c6"),
    CellTemplate("subsets", "Subsets · default", "Set", "FW_Subsets", "Complete power set.", "#287b95"),
    CellTemplate("subsets_exact", "Subsets · exact", "Set", "FW_Subsets_EXACT(2)", "Only subsets of the given size.", "#287b95"),
    CellTemplate("subsets_before", "Subsets · before", "Set", "FW_Subsets_BEFORE(2)", "Subset sizes below the parameter.", "#287b95"),
    CellTemplate("subsets_after", "Subsets · after", "Set", "FW_Subsets_AFTER(2)", "Subset sizes above the parameter.", "#287b95"),
    CellTemplate("subsets_range", "Subsets · range", "Set", "FW_Subsets_RANGE(2,3)", "Inclusive range of subset sizes.", "#287b95"),
    CellTemplate("subsets_given", "Subsets · given", "Set", "FW_Subsets_GIVEN(1,3)", "Only the listed subset sizes.", "#287b95"),
    CellTemplate("cartes", "Cartesian sheet", "Compose", "FW_Cartes(OTHER)", "Cross with all values of a named helper/source sheet.", "#7c3aed"),
    CellTemplate("cartes_first", "Cartesian baseline", "Compose", "FW_Cartes_first(OTHER)", "Cross with the other sheet's first value.", "#7c3aed"),
    CellTemplate("separator", "Separator", "Compose", "FW_Separator(SEPARATOR)", "Weave the first value of a helper sheet.", "#7c3aed"),
    CellTemplate("brace_1_1", "Join · 1:1", "Compose", "FW_(,,LEFT,,RIGHT,,,,1:1)", "Brace multiplicity 1:1.", "#7c3aed"),
    CellTemplate("brace_1_n", "Join · 1:N", "Compose", "FW_(,,LEFT,,RIGHT,,,,1:N)", "Brace multiplicity 1:N.", "#7c3aed"),
    CellTemplate("brace_m_1", "Join · M:1", "Compose", "FW_(,,LEFT,,RIGHT,,,,M:1)", "Brace multiplicity M:1.", "#7c3aed"),
    CellTemplate("brace_m_m", "Join · M:M", "Compose", "FW_(,,LEFT,,RIGHT,,,,M:M)", "Brace multiplicity M:M.", "#7c3aed"),
    CellTemplate("brace", "Join · M:N", "Compose", "FW_(,,LEFT,,RIGHT,,,,M:N)", "Only after two excluded operand rows.", "#7c3aed"),
    CellTemplate("brace_full_1_1", "Join full · 1:1", "Compose", "FW_(START,,LEFT,RELATION,RIGHT,,END,SEPARATOR,1:1)", "Nine-field join with start, relation, end and separator.", "#7c3aed"),
    CellTemplate("brace_full_1_n", "Join full · 1:N", "Compose", "FW_(START,,LEFT,RELATION,RIGHT,,END,SEPARATOR,1:N)", "Nine-field join with start, relation, end and separator.", "#7c3aed"),
    CellTemplate("brace_full_m_1", "Join full · M:1", "Compose", "FW_(START,,LEFT,RELATION,RIGHT,,END,SEPARATOR,M:1)", "Nine-field join with start, relation, end and separator.", "#7c3aed"),
    CellTemplate("brace_full_m_m", "Join full · M:M", "Compose", "FW_(START,,LEFT,RELATION,RIGHT,,END,SEPARATOR,M:M)", "Nine-field join with start, relation, end and separator.", "#7c3aed"),
    CellTemplate("brace_full_m_n", "Join full · M:N", "Compose", "FW_(START,,LEFT,RELATION,RIGHT,,END,SEPARATOR,M:N)", "Nine-field join with start, relation, end and separator.", "#7c3aed"),
    CellTemplate("brace_nested_left", "Join nested · left", "Compose", "FW_(,,FW_(),,RIGHT,,,,M:N)", "Use the most-recent prior brace target as the left operand.", "#7c3aed"),
    CellTemplate("brace_nested_right", "Join nested · right", "Compose", "FW_(,,LEFT,,FW_(),,,,M:N)", "Use the most-recent prior brace target as the right operand.", "#7c3aed"),
    CellTemplate("brace_nested_both", "Join nested · both", "Compose", "FW_(,,FW_(),,FW_(),,,,M:N)", "Use the two most-recent prior brace targets.", "#7c3aed"),
    CellTemplate("brace_grouped_left", "Join grouped · left", "Compose", "FW_(,,FW_()G,,RIGHT,,,,M:N)", "Use grouped output from the most-recent prior brace as left operand.", "#7c3aed"),
    CellTemplate("brace_grouped_right", "Join grouped · right", "Compose", "FW_(,,LEFT,,FW_()G,,,,M:N)", "Use grouped output from the most-recent prior brace as right operand.", "#7c3aed"),
    CellTemplate("brace_grouped_both", "Join grouped · both", "Compose", "FW_(,,FW_()G,,FW_()G,,,,M:N)", "Use grouped output from the two most-recent prior braces.", "#7c3aed"),
    CellTemplate("group", "FW_Group", "Transform", "FW_Group", "May occupy any directive column; its exact order is preserved.", "#536171"),
    CellTemplate("group_cleanup", "Group cleanup cell", "Transform", GROUP_CLEANUP, "Group plus ReplaceRE pipeline kept together in one XLSX cell.", "#536171"),
    CellTemplate("replace", "ReplaceRE", "Transform", 'FW_ReplaceRE("pattern", "replacement")', "Regex rewrite cell; double-click to edit.", "#536171"),
)
TEMPLATE_BY_ID = {template.id: template for template in TEMPLATES}


@dataclass
class GridRow:
    id: str
    target: str = ""
    directives: list[str] = field(default_factory=list)
    values: list[str] = field(default_factory=list)
    prefix: str = ""
    suffix: str = ""

    @property
    def is_used(self) -> bool:
        return bool(self.target.strip() or any(cell.strip() for cell in self.directives) or self.values or self.prefix or self.suffix)

    @property
    def is_blank(self) -> bool:
        return not self.is_used

    def directive(self, index: int) -> str:
        return self.directives[index] if 0 <= index < len(self.directives) else ""


@dataclass(frozen=True)
class GridIssue:
    severity: str
    message: str
    row_id: str = ""
    column: int | None = None


@dataclass
class GridProject:
    name: str
    rows: list[GridRow]
    directive_columns: int = 7
    auxiliary_sheets: dict[str, list[str]] = field(default_factory=dict)
    arguments: list[str] = field(default_factory=list)
    custom_verdicts: list[tuple[str, str]] = field(default_factory=list)
    run_once_source: str = ""
    _next_id: int = 1

    @classmethod
    def blank(cls, visible_rows: int = 12, directive_columns: int = 7) -> "GridProject":
        project = cls(name="my_fw_sequence", rows=[], directive_columns=directive_columns)
        starter = project.add_row("DIMENSION_1", ["baseline", "alternative"])
        starter.directives[0] = "FW_Combi(1)"
        project.add_blank_rows(max(visible_rows - 1, 0))
        return project

    def _new_row(self, target: str = "", values: Iterable[str] | None = None) -> GridRow:
        row = GridRow(
            id=f"row-{self._next_id}",
            target=target,
            directives=[""] * self.directive_columns,
            values=list(values or []),
        )
        self._next_id += 1
        return row

    def add_row(self, target: str = "", values: Iterable[str] | None = None) -> GridRow:
        row = self._new_row(target, values)
        self.rows.append(row)
        return row

    def add_blank_rows(self, count: int = 1) -> None:
        for _ in range(max(0, count)):
            self.rows.append(self._new_row())

    def add_directive_column(self) -> None:
        self.directive_columns += 1
        for row in self.rows:
            row.directives.append("")

    def row(self, row_id: str) -> GridRow | None:
        return next((row for row in self.rows if row.id == row_id), None)

    def set_cell(self, row_id: str, column: int, value: str) -> None:
        row = self.row(row_id)
        if row is None:
            return
        if column == 0:
            row.target = value
            return
        index = column - 1
        while index >= len(row.directives):
            self.add_directive_column()
        row.directives[index] = value

    def cell(self, row_id: str, column: int) -> str:
        row = self.row(row_id)
        if row is None:
            return ""
        return row.target if column == 0 else row.directive(column - 1)

    def move_row(self, row_id: str, delta: int) -> None:
        index = next((i for i, row in enumerate(self.rows) if row.id == row_id), -1)
        target = index + delta
        if index >= 0 and 0 <= target < len(self.rows):
            self.rows[index], self.rows[target] = self.rows[target], self.rows[index]

    @property
    def sequence_rows(self) -> list[GridRow]:
        last = max((index for index, row in enumerate(self.rows) if row.is_used), default=-1)
        return self.rows[:last + 1]

    def append_brace_trio(self) -> tuple[GridRow, GridRow, GridRow]:
        rows = self.sequence_rows
        occupied = {row.target for row in rows} | set(self.auxiliary_sheets)
        number = 1
        while any(name in occupied for name in (f"LEFT_{number}", f"RIGHT_{number}", f"JOIN_{number}")):
            number += 1
        insert_at = len(rows)
        while len(self.rows) < insert_at + 3:
            self.add_blank_rows()
        left, right, joined = self.rows[insert_at:insert_at + 3]
        left.target, left.values = f"LEFT_{number}", ["left_base", "left_alt"]
        right.target, right.values = f"RIGHT_{number}", ["right_base", "right_alt"]
        joined.target, joined.values = f"JOIN_{number}", [EMPTY]
        for operand in (left, right):
            operand.directives = [""] * self.directive_columns
            operand.directives[0:3] = ["FW_Exclude", "FW_Reuse", "FW_Combi(1)"]
        joined.directives = [""] * self.directive_columns
        joined.directives[0] = f"FW_(,,{left.target},,{right.target},,,,M:N)"
        return left, right, joined

    @staticmethod
    def _lines(row: GridRow) -> list[tuple[int, str]]:
        return [
            (column, line.strip())
            for column, cell in enumerate(row.directives, 1)
            for line in cell.splitlines()
            if line.strip()
        ]

    @classmethod
    def _positions(cls, row: GridRow, verb: str) -> list[int]:
        return [column for column, line in cls._lines(row) if line == verb]

    @staticmethod
    def _brace_cells(row: GridRow) -> list[tuple[int, str]]:
        return [
            (column, cell.strip())
            for column, cell in enumerate(row.directives, 1)
            if cell.strip().startswith("FW_(")
        ]

    @classmethod
    def _reuse_mode(cls, row: GridRow) -> str:
        if cls._positions(row, "FW_ReuseTableOnly"):
            return "FW_ReuseTableOnly"
        if cls._positions(row, "FW_Reuse"):
            return "FW_Reuse"
        return ""

    def validate(self) -> list[GridIssue]:
        issues: list[GridIssue] = []
        rows = self.sequence_rows
        if not rows:
            return [GridIssue("error", "FW_Seq needs at least one used row.")]

        targets = [row.target.strip() for row in rows]
        all_sheet_names = set(targets) | set(self.auxiliary_sheets)
        for index, row in enumerate(rows):
            excel_row = index + 1
            target = row.target.strip()
            if row.is_blank:
                issues.append(GridIssue("error", f"Row {excel_row} is blank inside the used FW_Seq range.", row.id))
                continue
            if not target:
                issues.append(GridIssue("error", f"A{excel_row}: target sheet name is required.", row.id, 0))
            elif len(target) > 31:
                issues.append(GridIssue("error", f"A{excel_row}: {target!r} exceeds Excel's 31-character sheet-name limit.", row.id, 0))
            elif INVALID_SHEET_CHARS.search(target):
                issues.append(GridIssue("error", f"A{excel_row}: {target!r} contains an invalid sheet-name character.", row.id, 0))
            elif target in CONTROL_SHEETS:
                issues.append(GridIssue("error", f"A{excel_row}: {target!r} is a reserved control sheet.", row.id, 0))
            if target and targets.count(target) > 1:
                issues.append(GridIssue("error", f"A{excel_row}: target {target!r} is duplicated.", row.id, 0))
            if not row.values:
                issues.append(GridIssue("error", f"{target or f'Row {excel_row}'} needs at least one source-data cell.", row.id, 0))
            for value in row.values:
                if value.startswith("FW_") and value != EMPTY:
                    issues.append(GridIssue("error", f"{target or f'Row {excel_row}'}: data value {value!r} would be parsed as a directive.", row.id, 0))
            if not any(cell.strip() for cell in row.directives):
                issues.append(GridIssue("error", f"Row {excel_row} has no directive cells.", row.id))

            optional = self._positions(row, "FW_Optional")
            exclude = self._positions(row, "FW_Exclude")
            reuse = self._positions(row, "FW_Reuse")
            reuse_table = self._positions(row, "FW_ReuseTableOnly")
            if optional and exclude:
                issues.append(GridIssue("error", f"Row {excel_row} cannot be both FW_Optional and FW_Exclude.", row.id))
            if reuse and reuse_table:
                issues.append(GridIssue("error", f"Row {excel_row} cannot contain both reuse modes.", row.id))
            occupied_positions = [column for column, cell in enumerate(row.directives, 1) if cell.strip()]
            for position in reuse + reuse_table:
                if not exclude:
                    issues.append(GridIssue("error", f"{get_column_letter(position + 1)}{excel_row}: reuse is allowed only on an FW_Exclude operand row.", row.id, position))
                else:
                    next_nonempty = next((item for item in occupied_positions if item > exclude[0]), None)
                    if position != next_nonempty:
                        issues.append(GridIssue("error", f"{get_column_letter(position + 1)}{excel_row}: reuse must be the next non-empty cell after FW_Exclude; blank cells between them are allowed.", row.id, position))

            braces = self._brace_cells(row)
            if len(braces) > 1:
                issues.append(GridIssue("error", f"Row {excel_row} contains more than one FW_(...) join cell.", row.id))
            for position, directive in braces:
                if index < 2:
                    issues.append(GridIssue("error", f"{get_column_letter(position + 1)}{excel_row}: a brace row needs two immediately preceding operand rows.", row.id, position))
                    continue
                left, right = rows[index - 2], rows[index - 1]
                if not self._positions(left, "FW_Exclude") or not self._positions(right, "FW_Exclude"):
                    issues.append(GridIssue("error", f"{get_column_letter(position + 1)}{excel_row}: FW_(...) is valid only directly after two FW_Exclude rows.", row.id, position))
                match = re.fullmatch(r"FW_\((.*)\)", directive, flags=re.DOTALL)
                fields = [part.strip() for part in match.group(1).split(",")] if match else []
                if len(fields) != 9:
                    issues.append(GridIssue("error", f"{get_column_letter(position + 1)}{excel_row}: FW_(...) needs exactly nine comma-separated fields.", row.id, position))
                elif any(
                    operand not in {target, "FW_()", "FW_()G"}
                    for operand, target in (
                        (fields[2], left.target.strip()),
                        (fields[4], right.target.strip()),
                    )
                ):
                    issues.append(GridIssue("error", f"{get_column_letter(position + 1)}{excel_row}: brace operands must be the two preceding targets (or Core nested markers FW_()/FW_()G), {left.target!r} then {right.target!r}.", row.id, position))
                modes = (self._reuse_mode(left), self._reuse_mode(right))
                if bool(modes[0]) != bool(modes[1]) or (modes[0] and modes[0] != modes[1]):
                    issues.append(GridIssue("error", f"Rows {excel_row - 2}–{excel_row - 1}: both brace operands must use the same reuse mode, or neither.", row.id))

            for pattern in (r"FW_Separator\(([^)]+)\)", r"FW_Cartes(?:_first)?\(([^)]+)\)"):
                for _, line in self._lines(row):
                    for reference in re.findall(pattern, line):
                        if reference not in all_sheet_names:
                            issues.append(GridIssue("error", f"Row {excel_row} references missing sheet {reference!r}.", row.id))

        brace_indexes = {index for index, row in enumerate(rows) if self._brace_cells(row)}
        for index, row in enumerate(rows):
            if not self._reuse_mode(row):
                continue
            if not any(index in (brace_index - 2, brace_index - 1) for brace_index in brace_indexes):
                issues.append(GridIssue("error", f"Row {index + 1}: {self._reuse_mode(row)} is allowed only on one of the two rows immediately before FW_(...).", row.id))

        for sheet_name in self.auxiliary_sheets:
            if not sheet_name or len(sheet_name) > 31 or INVALID_SHEET_CHARS.search(sheet_name) or sheet_name in CONTROL_SHEETS:
                issues.append(GridIssue("error", f"Auxiliary sheet name {sheet_name!r} is invalid or reserved."))
            if sheet_name in targets:
                issues.append(GridIssue("error", f"Auxiliary sheet {sheet_name!r} duplicates a sequence target."))

        if self.run_once_source:
            missing = [
                label
                for pattern, label in (
                    (r"class\s+RunMeFirstOnce", "class RunMeFirstOnce"),
                    (r"public\s+static\s+String\s+FW_ARGS", "public static String FW_ARGS"),
                    (r"FW_ARGS\s*=", "FW_ARGS assignment"),
                )
                if not re.search(pattern, self.run_once_source)
            ]
            if missing:
                issues.append(GridIssue("warning", "FW_RunMeFirstOnce A1 is missing Core's expected " + ", ".join(missing) + "."))
            if len(self.run_once_source) > 32_767:
                issues.append(GridIssue("error", "FW_RunMeFirstOnce A1 exceeds Excel's 32,767-character cell limit."))

        for row_index, argument in enumerate(self.arguments, 1):
            if len(str(argument)) > 32_767:
                issues.append(GridIssue("error", f"FW_Arguments A{row_index} exceeds Excel's 32,767-character cell limit."))

        seen_verdict_codes: set[int] = set()
        for row_index, (code, message) in enumerate(self.custom_verdicts, 1):
            code_text = str(code).strip()
            if not re.fullmatch(r"[+-]?\d+", code_text):
                issues.append(GridIssue("error", f"FW_CUSTOM_VAR A{row_index} must be a Java integer, got {code!r}."))
            else:
                parsed_code = int(code_text)
                if not -(2**31) <= parsed_code < 2**31:
                    issues.append(GridIssue("error", f"FW_CUSTOM_VAR A{row_index} is outside Java Integer range."))
                elif parsed_code in seen_verdict_codes:
                    issues.append(GridIssue("error", f"FW_CUSTOM_VAR A{row_index} duplicates verdict code {parsed_code}."))
                seen_verdict_codes.add(parsed_code)
            if len(str(message)) > 32_767:
                issues.append(GridIssue("error", f"FW_CUSTOM_VAR B{row_index} exceeds Excel's 32,767-character cell limit."))
        return issues

    def to_xlsx_bytes(self) -> bytes:
        errors = [issue for issue in self.validate() if issue.severity == "error"]
        if errors:
            raise ValueError("Cannot export invalid grid: " + "; ".join(issue.message for issue in errors[:5]))

        workbook = Workbook()
        workbook.remove(workbook.active)
        seq = workbook.create_sheet("FW_Seq")
        names = workbook.create_sheet("FW_SheetNames")
        run_once = workbook.create_sheet("FW_RunMeFirstOnce")
        arguments = workbook.create_sheet("FW_Arguments")
        verdicts = workbook.create_sheet("FW_CUSTOM_VAR")
        info = workbook.create_sheet("FW_Info")

        rows = self.sequence_rows
        for row_index, row in enumerate(rows, 1):
            seq.cell(row_index, 1, row.target.strip())
            info.cell(row_index, 1, row.target.strip())
            for index, directive in enumerate(row.directives, 2):
                if directive.strip():
                    seq.cell(row_index, index, directive.strip())
                    info.cell(row_index, index, directive.strip())
            names.cell(row_index, 1, row.target.strip())
            names.cell(row_index, 2, row.prefix or EMPTY)
            names.cell(row_index, 3, row.suffix or EMPTY)

        if self.run_once_source:
            run_once["A1"] = self.run_once_source
        for row_index, argument in enumerate(self.arguments, 1):
            arguments.cell(row_index, 1, argument or EMPTY)
        for row_index, (code, message) in enumerate(self.custom_verdicts, 1):
            verdicts.cell(row_index, 1, int(str(code).strip()))
            verdicts.cell(row_index, 2, message or EMPTY)

        written = set(CONTROL_SHEETS)
        for row in rows:
            ws = workbook.create_sheet(row.target.strip())
            for value_index, value in enumerate(row.values, 1):
                ws.cell(value_index, 1, value)
            written.add(row.target.strip())
        for sheet_name, values in self.auxiliary_sheets.items():
            if sheet_name in written:
                continue
            ws = workbook.create_sheet(sheet_name)
            for value_index, value in enumerate(values, 1):
                ws.cell(value_index, 1, value)
            written.add(sheet_name)

        header_fill = PatternFill("solid", fgColor="E8F3F1")
        for ws in (seq, names, info):
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
            for column in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(column)].width = 22
            for row_cells in ws.iter_rows():
                for cell in row_cells:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
        workbook.active = workbook.sheetnames.index("FW_Info")
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()


def project_from_analysis(analysis: WorkbookAnalysis, trailing_rows: int = 8) -> GridProject:
    """Import an analyzed workbook only when the user explicitly chooses it."""
    width = max((len(step.directives) for step in analysis.sequence_steps), default=7)
    project = GridProject(
        name=re.sub(r"[^A-Za-z0-9_.-]+", "_", analysis.source_name.rsplit(".", 1)[0]) or "fw_sequence",
        rows=[],
        directive_columns=max(7, width),
        arguments=list(analysis.runtime.arguments),
        custom_verdicts=list(analysis.runtime.custom_verdicts),
        run_once_source=analysis.runtime.run_once_source,
    )
    target_names = {step.sheet_name for step in analysis.sequence_steps}
    for step in analysis.sequence_steps:
        sheet = analysis.sheet_by_name[step.sheet_name]
        row = project.add_row(step.sheet_name, [cell.value for cell in sheet.values])
        row.directives[:len(step.directives)] = list(step.directives)
        row.prefix = "" if step.prefix == EMPTY else step.prefix
        row.suffix = "" if step.suffix == EMPTY else step.suffix
    project.auxiliary_sheets = {
        sheet.name: [cell.value for cell in sheet.values]
        for sheet in analysis.sheets
        if sheet.role != "control" and sheet.name not in target_names
    }
    project.add_blank_rows(trailing_rows)
    return project
