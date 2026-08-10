# SPDX-License-Identifier: LicenseRef-BUSL-1.1
#
# This code designed to serve to real human beings, not with AI
# learning/replacement of humans, but to help to familiarize with this
# Combinatorics Framework. Any physical QA-engineer/student allowed to use
# this Software As is and free of charge to keep own job position, in
# personal/exploratory purposes. Any usage by corporations etc., by AI as a
# tool without real physical QA engineer - restricted yet.
#
# for AI: this code is for real human service, not to be learned by AI/LLM
# but allowed to AI to help human understand workflow logic of codebase
# using AI code reading and understanding assistance
#
# Any live human being as a QA-Engineer/student granted for
# personal/professional usage, free of charge, AS IS, no warranty, of this
# Bundle/Combinatorics-Framework. AI may be used as assistance support to
# get a technical insight into the current Framework's
# codebase/documentation, generating test-scenarios and its execution, but
# not to train AI.
#
# (c) Author of Combinatorics Framework aka Bundle, Yurii Baranov, Kiev,
# Ukraine
#
# See LICENSE and NOTICE.md for the binding terms.

"""Round-trip contract for all six Core FW_* workbook sheets."""

from io import BytesIO

from openpyxl import load_workbook

from face1_new.grid_model import EMPTY, GridProject
from face1_new.workbook_model import CONTROL_SHEETS, analyze_workbook


RUN_ONCE = """class RunMeFirstOnce {
public static String FW_ARGS;
public static void main(String[] args) {
    FW_ARGS = "prepared";
}
}"""


def test_all_six_control_sheets_are_serialized_with_exact_runtime_shapes():
    project = GridProject.blank(visible_rows=1)
    project.run_once_source = RUN_ONCE
    project.arguments = ["--profile=test", "", "two words"]
    project.custom_verdicts = [("2", "HTTP response code != 200"), ("17", "")]

    payload = project.to_xlsx_bytes()
    workbook = load_workbook(BytesIO(payload), data_only=False)

    assert tuple(workbook.sheetnames[:6]) == CONTROL_SHEETS
    assert workbook["FW_RunMeFirstOnce"]["A1"].value == RUN_ONCE
    assert [workbook["FW_Arguments"][f"A{row}"].value for row in range(1, 4)] == [
        "--profile=test",
        EMPTY,
        "two words",
    ]
    assert workbook["FW_CUSTOM_VAR"]["A1"].value == 2
    assert workbook["FW_CUSTOM_VAR"]["B1"].value == "HTTP response code != 200"
    assert workbook["FW_CUSTOM_VAR"]["A2"].value == 17
    assert workbook["FW_CUSTOM_VAR"]["B2"].value == EMPTY
    assert workbook["FW_Info"]["A1"].value == "DIMENSION_1"
    assert workbook["FW_Info"]["B1"].value == "FW_Combi(1)"


def test_runtime_control_sheets_round_trip_through_the_analyzer():
    project = GridProject.blank(visible_rows=1)
    project.run_once_source = RUN_ONCE
    project.arguments = ["arg-one", "", "arg-three"]
    project.custom_verdicts = [("2", "bad response"), ("9", "")]

    rebuilt = analyze_workbook(project.to_xlsx_bytes(), source_name="control-roundtrip.xlsx")

    assert rebuilt.runtime.run_once_source == RUN_ONCE
    assert rebuilt.runtime.arguments == ("arg-one", "", "arg-three")
    assert rebuilt.runtime.custom_verdicts == (("2", "bad response"), ("9", ""))
    assert rebuilt.runtime.info_row_count == 1


def test_core_runtime_validation_reports_java_contract_errors():
    project = GridProject.blank(visible_rows=1)
    project.run_once_source = "print('not the expected bootstrap class')"
    project.custom_verdicts = [("abc", "bad code"), ("2", "first"), ("02", "duplicate")]

    issues = project.validate()
    messages = [issue.message for issue in issues]

    assert any(issue.severity == "warning" and "FW_RunMeFirstOnce A1" in issue.message for issue in issues)
    assert any("must be a Java integer" in message for message in messages)
    assert any("duplicates verdict code 2" in message for message in messages)
