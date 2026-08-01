r"""fwgen — headless, data-driven generator of Core input-test workbooks.

A synthesis of two predecessors:

  * v25 "Pro Data Engine" (Pictures/v25_xlsx_to_json25052026/) — a PyQt6 GUI
    wrapping a pure-itertools generation engine, OPT-IN N-wise covering-array
    reduction (greedy + optimal set-cover; default = exhaustive), chunking, and
    XLSX/JSON output. Strength: real reduction + dual output. Weakness: locked
    inside a Qt worker thread (can't script / CI / test headlessly).

  * FW_scenario_inputs/gen_scenarios.py (today) — code-verified FW_ control-sheet
    formats, compact "cartesian-leaves" wiring (the *Core* does the expansion),
    k=1 directed ablation, validation + README. Weakness: scenario data was
    HARDCODED in the Python.

This module keeps the strengths and drops the weaknesses:

  * DATA-HARDCODE-FREE. There is ZERO scenario data in this file. Every scenario
    lives in an external spec (specs/*.toml|*.json|*.yaml). The code only knows
    how to READ a spec and BUILD workbooks.
  * SMART AUTO-INFERENCE ("let it pick the right variants"): baseline = first
    value of each slot; Pareto goal direction inferred from the metric-name
    lexicon; reduction strength auto-chosen against a compute budget. Exhaustive
    by DEFAULT — the user opts into reduction (the author's stated philosophy:
    don't cap the space, let the operator self-limit with the exponential wall
    in mind).
  * CORE-FAITHFUL. FW_ sheet formats verified against the Core source
    (WorkbookParser / SeqParser / MainRefactored):
      - any sheet whose name starts with "FW_" is a control sheet; data sheets
        must NOT start with "FW_", and data CELLS must not start with "FW_"
        (else parsed as cell-directives) — our cells begin with a space;
      - FW_Exclude/FW_Heading set a SHEET aside from the final cartesian (into
        the Core's exclude map) for a brace-join FW_(...) to consume — they are
        SHEET-level, NOT a "drop this combination" row filter; FW_Optional routes
        a sheet to fw_opt<i>;
      - a brace FW_(...) JOINS the two operand RESULT TABLES named at positions 2 & 4
        — i.e. the combination-RESULTS (fw2_<sheet>/fw_<sheet>) that EARLIER FW_Seq
        rows produced for those sheets, NOT the raw value sheets. FW_Seq is a dataflow:
        each row makes a result table; the brace joins two prior ones by relation +
        multiplicity (1:1/1:N/M:1/M:M/M:N) and writes a NEW result table. Operands MUST
        be FW_Exclude'd so they appear ONLY via the join (Core: SeqParser.parseExcludedKeys
        + BraceOperationHandler.{processOneToFormula,processManyToFormula,cleanupExcludedTables});
        NOTE: author seq_extra as a TOP-LEVEL toml key (before any [[slots]]/[[goals]]
        table) or it binds to the last table and is silently dropped;
      - FW_Reuse / FW_ReuseTableOnly preserve a brace operand's table / rows from
        the join's cleanup (Reuse=keep table, ReuseTableOnly=keep rows; nested
        operands need both). On a PLAIN (non-operand) slot FW_Reuse is a harmless
        no-op — a slot is in the cartesian by simply NOT being excluded;
      - a lone FW_Combi(1) auto-promotes to dual;
      - FW_SheetNames: col0=sheet, col1=prefix, col2=ending; FW_EMPTY_STRING→"";
      - FW_CUSTOM_VAR: col0 MUST parse as int, col1=message;
      - FW_Info is inert (stored, never parsed) — pure documentation.
"""
from __future__ import annotations

import itertools
import json
import logging
import math
import re
from concurrent.futures import ProcessPoolExecutor
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path
from typing import Iterable, Iterator, Optional, Sequence

import openpyxl
from openpyxl import Workbook

from xlsx_autofit import autofit_workbook   # ported from v25 xlsx_view_resizer

log = logging.getLogger("fwgen")

# ===========================================================================
# 1. Shared text conventions  (lifted from v25 combiner_core, trimmed)
# ===========================================================================
# Two-character sequences in data (backslash + letter), NOT control chars.
LITERAL_TAB = "\\t"          # in-token name/tag separator
LITERAL_NEWLINE = "\\n"      # intra-token newline (multi-line cell, e.g. FW_Group)


def cell_for_xlsx(text: str) -> str:
    r"""Render a stored cell value for XLSX: literal \n -> real in-cell newline,
    literal \t -> real tab. Lets multi-line directive values (FW_Group...) show
    as one multi-line cell, exactly as the Core's source workbooks do."""
    return text.replace(LITERAL_NEWLINE, "\n").replace(LITERAL_TAB, "\t")


# ===========================================================================
# 2. Spec model + loaders + smart inference
# ===========================================================================

# Lexicon for auto-inferring Pareto goal direction (mirrors the Analyzer's
# AutoAnalysisPlanner: cost/latency/error -> MIN, throughput/accuracy -> MAX).
_MIN_TOKENS = ("cost", "price", "latency", "lat", "error", "err", "loss",
               "size", "time", "duration", "ms", "wait", "far", "frr",
               "cycles", "length", "p95", "p99", "memory", "ram", "bytes")
_MAX_TOKENS = ("throughput", "tps", "rps", "accuracy", "acc", "f1", "score",
               "coverage", "recall", "precision", "map", "fps", "relevance",
               "bypass", "crash", "success", "quality", "reward")


def infer_direction(key: str) -> str:
    """min | max — inferred from the metric name; defaults to 'max' if unknown."""
    k = key.lower()
    if any(t in k for t in _MIN_TOKENS):
        return "min"
    if any(t in k for t in _MAX_TOKENS):
        return "max"
    return "max"


# ---------------------------------------------------------------------------
# Core FW_Seq combinatorial grammar — verified against the Core source
# (SeqParser.isComboRuleVerb/isAlgorithmDirective/FW_BRACE_VALID,
#  SheetWorker.resolveAlgoType/resolveM/buildGenerator,
#  CombinatorialGenerator.SubsetMode). This lets a spec choose ANY combination
# rule per slot — not just FW_Combi(1) — so a generated workbook can exercise
# the engine's whole vocabulary. DEFAULT_VERB keeps the historical behaviour.
# ---------------------------------------------------------------------------
DEFAULT_VERB = "FW_Combi(1)"

# Unary verbs (operate on ONE sheet's rows). Arg forms: k=int | all|full (every
# size, engine m=-2) | size (= list length); bare or () => k=1. R-variants first
# so e.g. "FW_PermutR(2)" can't be mis-read as "FW_Permut".
_VERB_RE = re.compile(
    r"FW_CombiR(\(\d+\)|\(all\)|\(full\)|\(\))?"                          # multicombinations
    r"|FW_Combi(\(\d+\)|\(all\)|\(full\)|\(size\)|\(\))?"                 # combinations C(n,k)
    r"|FW_PermutR(\(\d+\)|\(all\)|\(full\)|\(\))?"                        # permutations w/ repetition n^k
    r"|FW_Permut(\(\d+\)|\(\)|\(PermutationGenerator\.TreatDuplicatesAs\.IDENTICAL\))?"  # perms / k-perms / multiset
    r"|FW_Subsets(_(BEFORE|AFTER|EXACT|RANGE|GIVEN)\(\d+(,\d+)*\)|\(\d+\))?"  # powerset 2^n / bounded
    r"|FW_Cartes(_first)?\([A-Za-z0-9_]+\)"                              # cartesian with another sheet
    r"|FW_Separator\([A-Za-z0-9_]+\)")                                   # interleave a separator sheet
# Brace cross-sheet joiner (mirrors SeqParser.FW_BRACE_VALID): 9 fields / 8
# commas, fields 1 & 5 empty, multiplicity ∈ {1:1,1:N,M:1,M:M,M:N}.
# Either operand may be Core's nested marker: FW_() consumes the most-recent
# prior brace result as rows; FW_()G consumes it as one grouped row.
_BRACE_OPERAND = r"(?:[A-Za-z0-9_]+|FW_\(\)G?)"
_BRACE_RE = re.compile(
    rf"FW_\([A-Za-z0-9_]*,,{_BRACE_OPERAND},[A-Za-z0-9_]*,"
    rf"{_BRACE_OPERAND},,[A-Za-z0-9_]*,[A-Za-z0-9_]*,"
    r"(1:[1Nn]|[Mm]:[1MmNn])\)")
_FLAG_RE = re.compile(
    r"FW_Optional|FW_Exclude|FW_Heading|FW_LastInQueue|FW_Reuse|FW_ReuseTableOnly|FW_Concatenator=.*")
# operands a verb references as OTHER sheets (must exist), for validation.
_OPERAND_RE = re.compile(r"FW_(?:Cartes(?:_first)?|Separator)\(([A-Za-z0-9_]+)\)")


def _first_line(tok: str) -> str:
    s = str(tok).strip()
    return s.splitlines()[0].strip() if s else ""


def is_core_verb(tok: str) -> bool:
    """True if `tok` is a Core combinatorial verb / brace joiner the engine runs.
    FW_Group is multi-line (FW_Group + embedded FW_ReplaceRE(...)): its first line
    is 'FW_Group'."""
    first = _first_line(tok)
    if first.startswith("FW_Group"):
        return True
    return bool(_VERB_RE.fullmatch(first) or _BRACE_RE.fullmatch(first))


def is_core_flag(tok: str) -> bool:
    return bool(_FLAG_RE.fullmatch(str(tok).strip()))


# STEP 36: domain-level authoring aliases. A slot may be authored with a plain-
# English `alias` INSTEAD of a raw `verb`; the compiler (`compile_alias`) lowers
# it to the SAME advanced FW verb (+ flags) an expert would have written. Aliases
# are opt-in sugar — the expert `verb`/flags/braces path is untouched, and an
# alias and its equivalent expert spec MUST produce the identical workbook/count.
# Argument style mirrors the verbs themselves: an optional "(k)" suffix.
ALIAS_NAMES = frozenset({
    "choose_one", "choose_k", "permute", "feature_subset", "optional_action",
})
_ALIAS_RE = re.compile(r"([a-z_]+)(?:\((\d+)\))?$")


def is_alias(tok: str) -> bool:
    """True if `tok` is a recognized domain-level alias (with or without a (k) arg)."""
    m = _ALIAS_RE.fullmatch(_first_line(tok))
    return bool(m) and m.group(1) in ALIAS_NAMES


def compile_alias(tok: str) -> tuple[str, tuple[str, ...]]:
    """Lower a domain alias to ``(verb, extra_flags)`` — the advanced FW
    representation. Raises ValueError on an unknown alias or a missing/forbidden
    ``(k)`` argument. Mapping (each is the literal expert verb the engine runs):

    - ``choose_one``        -> ``FW_Combi(1)``          (pick exactly one value)
    - ``choose_k(k)``       -> ``FW_Combi(k)``          (pick exactly k, C(n,k))
    - ``permute`` / ``(k)`` -> ``FW_Permut`` / ``FW_Permut(k)`` (orderings / k-perms)
    - ``feature_subset``    -> ``FW_Subsets``           (powerset 2^n)
    - ``feature_subset(k)`` -> ``FW_Subsets(k)``        (bounded powerset)
    - ``optional_action``   -> ``FW_Combi(1)`` + flag ``FW_Optional`` (present-or-absent)
    """
    m = _ALIAS_RE.fullmatch(_first_line(tok))
    if not m or m.group(1) not in ALIAS_NAMES:
        raise ValueError(f"unknown authoring alias {tok!r} "
                         f"(supported: {', '.join(sorted(ALIAS_NAMES))})")
    name, arg = m.group(1), m.group(2)
    k = int(arg) if arg is not None else None

    def _no_arg():
        if k is not None:
            raise ValueError(f"alias {name!r} takes no (k) argument, got {tok!r}")

    if name == "choose_one":
        _no_arg()
        return "FW_Combi(1)", ()
    if name == "choose_k":
        if k is None:
            raise ValueError(f"alias 'choose_k' requires a (k) argument, e.g. choose_k(2); got {tok!r}")
        if k < 1:
            raise ValueError(f"alias 'choose_k(k)' needs k >= 1, got {tok!r}")
        return f"FW_Combi({k})", ()
    if name == "permute":
        return (f"FW_Permut({k})" if k is not None else "FW_Permut"), ()
    if name == "feature_subset":
        return (f"FW_Subsets({k})" if k is not None else "FW_Subsets"), ()
    if name == "optional_action":
        _no_arg()
        return "FW_Combi(1)", ("FW_Optional",)
    raise ValueError(f"unhandled alias {tok!r}")  # pragma: no cover


def verb_sheet_operands(tok: str) -> list[str]:
    """Other-sheet names a verb/joiner references (Cartes/Separator operand, or
    brace operands Ei,Ej) — for cross-checking that those sheets exist."""
    first = _first_line(tok)
    ops = _OPERAND_RE.findall(first)
    if _BRACE_RE.fullmatch(first):
        inner = first[first.index("(") + 1:first.rindex(")")]
        parts = inner.split(",")
        # 9 fields: FW_(start,_,E1,rel,E2,_,end,sep,mult) — every SHEET-name field must exist:
        # start(0), E1(2), rel(3), E2(4), end(6), sep(7). (1,5 are empty; 8 is the multiplicity.)
        ops += [
            p
            for i, p in enumerate(parts)
            if i in (0, 2, 3, 4, 6, 7)
            and p
            and p not in {"FW_()", "FW_()G"}
        ]
    return ops


def verb_output_count(verb: str, n: int, other_n: int = 0) -> int:
    """Estimate how many rows the Core's engine emits for ONE slot of `n` values
    under `verb` (other_n = operand size for Cartes). Mirrors CombinatorialGenerator
    / SheetWorker.resolveM. Best-effort — lets us PREDICT the fw_final size and warn
    before emitting, so the inner combined data stays small (no explosion)."""
    v = _first_line(verb)

    def _arg(default=1):
        m = re.search(r"\((\d+)\)", v)
        return int(m.group(1)) if m else default

    if v.startswith(("FW_Group", "FW_Separator")):
        return max(1, n)                                    # ~row-preserving transforms
    if v.startswith("FW_Cartes"):
        return max(1, n) * max(1, other_n or 1)
    if v.startswith("FW_CombiR"):
        if re.search(r"\((all|full)\)", v):
            return sum(math.comb(n + k - 1, k) for k in range(1, n + 1)) or 1
        k = _arg()
        return math.comb(n + k - 1, k) if k >= 0 and n > 0 else 1
    if v.startswith("FW_Combi"):
        if re.search(r"\((all|full)\)", v):
            return (2 ** n - 1) or 1
        if re.search(r"\(size\)", v):
            return 1
        k = _arg()
        return math.comb(n, k) if 0 <= k <= n else 1
    if v.startswith("FW_PermutR"):
        if re.search(r"\((all|full)\)", v):
            return sum(n ** k for k in range(1, n + 1)) or 1
        return n ** _arg()
    if v.startswith("FW_Permut"):
        if re.search(r"\(\d+\)", v):
            k = _arg()
            return math.perm(n, k) if 0 <= k <= n else 1
        return math.factorial(n) or 1                       # incl. multiset (Core streams all n!)
    if v.startswith("FW_Subsets_"):
        sizes = [int(x) for x in re.findall(r"\d+", v)]
        mode = re.search(r"FW_Subsets_(\w+)", v).group(1).upper()
        if mode == "EXACT":
            return math.comb(n, sizes[0]) if sizes and sizes[0] <= n else 1
        if mode == "BEFORE":
            return sum(math.comb(n, k) for k in range(0, sizes[0])) if sizes else 1
        if mode == "AFTER":
            return sum(math.comb(n, k) for k in range(sizes[0] + 1, n + 1)) if sizes else 1
        if mode == "RANGE":
            return sum(math.comb(n, k) for k in range(sizes[0], sizes[1] + 1)) if len(sizes) >= 2 else 1
        if mode == "GIVEN":
            return sum(math.comb(n, k) for k in sizes if 0 <= k <= n) or 1
    if v.startswith("FW_Subsets"):
        return 2 ** n
    return max(1, n)                                        # unknown → assume row-preserving


# Verbs whose row count is a closed-form function of (n, k) — the formula sees
# every input it needs, so the prediction is EXACT, not a best-effort guess.
_EXACT_FORMULA_PREFIXES = (
    "FW_CombiR", "FW_Combi", "FW_PermutR", "FW_Permut", "FW_Subsets",
)


def verb_cardinality(verb: str, n: int, other_n: int = 0) -> CardinalityEstimate:
    """Classify ONE slot's verb output by confidence (STEP 9), reusing
    :func:`verb_output_count` for the arithmetic but being honest about which
    verbs give a provable number and which only a best-effort one.

    - First-order combinatorial verbs (Combi/CombiR/Permut/PermutR/Subsets*):
      EXACT — closed-form C(n,k)/n!/2^n etc. over fully-known operands.
    - FW_Separator: EXACT — deterministic weave, row count == operand count.
    - FW_Cartes(OTHER): BOUNDED — `other_n` is the OPERAND SHEET's declared
      size, not necessarily that sheet's actual Core result-table row count
      (which can differ after its own verb/exclusion/join effects), so the
      product is an upper bound, not a proof.
    - FW_Group: BOUNDED — row-preserving EXCEPT that its `group_replace`
      FW_ReplaceRE rewrites can make two distinct rows collide post-rewrite;
      the engine's dedup behavior isn't statically known, so [1, n] is the
      provable range (see plan STEP 9 action 4: never call this exact).
    - anything else (verb the model doesn't recognize): UNKNOWN.
    """
    v = _first_line(verb)
    value = verb_output_count(verb, n, other_n)

    if v.startswith(_EXACT_FORMULA_PREFIXES):
        return CardinalityEstimate.exact(value, formula=f"{v} over n={n}",
                                          assumptions=("operand value count `n` is the declared slot size",))
    if v.startswith("FW_Separator"):
        return CardinalityEstimate.exact(value, formula=f"max(1, n) — {v} weaves without changing row count",
                                          assumptions=("operand sheet provides at least one value to weave",))
    if v.startswith("FW_Cartes"):
        lower = max(1, n, other_n or 1)
        return CardinalityEstimate.bounded(
            lower=lower, upper=value, value=value,
            formula=f"n × other_n = {n} × {other_n} (upper bound)",
            reasons=(f"{v} joins this slot's {n} values against the operand SHEET's declared "
                     f"{other_n} values, but the operand's actual Core result-table size can differ "
                     f"after its own verb/exclusion/brace effects — the product is an upper bound only",),
            assumptions=("operand sheet row count approximates its post-verb result-table size",))
    if v.startswith("FW_Group"):
        return CardinalityEstimate.bounded(
            lower=1, upper=max(1, n),
            formula=f"<= max(1, n) = {max(1, n)} (row-preserving transform, upper bound)",
            reasons=(f"{v}'s group_replace FW_ReplaceRE rewrites operate per-row and cannot create new "
                     f"rows, but CAN make distinct rows collide post-rewrite — whether the engine then "
                     f"dedupes is not statically known, so only [1, n] is provable",),
            assumptions=("no information about the engine's post-rewrite deduplication behavior",))
    return CardinalityEstimate.unknown(
        formula=f"verb {v!r} not recognized by the cardinality model",
        reasons=(f"falls back to verb_output_count's row-preserving guess ({value}), which is "
                 f"unverified for this verb — treat as informational only, not a bound",))


def brace_cardinality(expr: str) -> CardinalityEstimate:
    """Classify a brace-joiner FW_(...) row from `seq_extra` (STEP 9 action 4).

    A brace JOINS the result tables of its two operand slots — a second-order
    effect that depends on each operand's actual post-verb row count, which
    `estimate_core_combos`/`verb_cardinality` cannot see (they only know
    declared slot sizes). Per the plan, never present this as an approximate
    number dressed up as exact: it is UNKNOWN, with the reason spelled out.
    """
    first = _first_line(expr)
    return CardinalityEstimate.unknown(
        formula=f"brace {first!r} joins two operand result tables",
        reasons=(f"brace joiner {first!r} combines the RESULT TABLES of its operand slots "
                 f"(not their raw declared sizes); that join's row count depends on each "
                 f"operand's actual post-verb/post-exclusion Core output, which is only known "
                 f"by running Core — no statically sound bound exists here",))


def estimate_core_combos(spec: "Spec") -> int:
    """Approximate fw_final row count = product over the MANDATORY (non-excluded,
    non-optional) slots of each slot's verb output. Cartes operands sized via the
    referenced sheet. Approximate (ignores brace/join effects) — for explosion
    awareness, NOT an exact count."""
    sheet_n = {s.sheet: len(s.values) for s in spec.slots}
    prod = 1
    for s in spec.slots:
        if any(f in s.flags for f in ("FW_Exclude", "FW_Heading", "FW_Optional")):
            continue
        ops = verb_sheet_operands(s.verb)
        other = sheet_n.get(ops[0], 0) if ops else 0
        prod *= max(1, verb_output_count(s.verb, len(s.values), other))
    return prod


def optional_multiplier_cardinality(spec: "Spec") -> CardinalityEstimate:
    """EXACT cardinality of the FW_Optional expansion factor (STEP 9 action 3).

    Each FW_Optional slot independently contributes "any of its values, OR
    absent" — a closed-form `len(values) + 1` per slot, multiplied together.
    Mirrors bundle/cli.py's `optional_factor` computation, but as a reusable,
    confidence-tagged Spec-level API (no behavioral change to that call site)."""
    opt_slots = [s for s in spec.slots if "FW_Optional" in s.flags]
    if not opt_slots:
        return CardinalityEstimate.exact(1, formula="no FW_Optional slots -> factor 1")
    factor = 1
    terms = []
    for s in opt_slots:
        n = len(s.values) + 1                                # its values, OR absent
        factor *= n
        terms.append(f"({len(s.values)}+1 absent)[{s.sheet}]")
    return CardinalityEstimate.exact(
        factor, formula=" × ".join(terms),
        assumptions=("each FW_Optional slot independently toggles present/absent "
                     "(no cross-slot dependency between optional choices)",))


def _combine_product(estimates: "Sequence[CardinalityEstimate]", *, formula: str,
                      reasons: "Sequence[str]" = ()) -> CardinalityEstimate:
    """Multiply confidence-tagged estimates, propagating the weakest tier.

    EXACT × EXACT × ... -> EXACT (every factor is provably the stated number).
    Any BOUNDED factor (none UNKNOWN) -> BOUNDED, multiplying lower/upper/value
    componentwise — still a provable range, just no longer a single proven point.
    Any UNKNOWN factor -> UNKNOWN: a join with an unknowable operand makes the
    whole product unknowable too (no silently inventing a number or a bound)."""
    if not estimates:
        return CardinalityEstimate.exact(1, formula=f"{formula} (empty product = 1)")
    if any(e.mode == CardinalityMode.UNKNOWN for e in estimates):
        culprits = [e.formula for e in estimates if e.mode == CardinalityMode.UNKNOWN]
        return CardinalityEstimate.unknown(
            formula=formula,
            reasons=tuple(reasons) + (f"product includes UNKNOWN factor(s): {culprits}",))
    value = lower = upper = 1
    for e in estimates:
        value *= e.value
        lower *= e.lower
        upper *= e.upper
    if all(e.mode == CardinalityMode.EXACT for e in estimates):
        return CardinalityEstimate.exact(value, formula=formula)
    return CardinalityEstimate.bounded(lower=lower, upper=upper, value=value, formula=formula,
                                        reasons=tuple(reasons) + (
                                            "product includes BOUNDED factor(s); range "
                                            "multiplies componentwise (lower×lower, upper×upper)",))


@dataclass(frozen=True)
class SpecCardinalityPlan:
    """Full per-spec cardinality breakdown with confidence at every stage
    (STEP 9 action 5): raw declared sizes -> per-slot Core rows -> mandatory
    product -> post-sieve estimate -> optional multiplier -> final count.
    Each stage keeps its own :class:`CardinalityEstimate` (formula + bounds +
    reasoning) instead of collapsing everything to one bare integer."""
    raw_values: "dict[str, int]"
    per_slot: "dict[str, CardinalityEstimate]"
    mandatory: CardinalityEstimate
    post_sieve: CardinalityEstimate
    optional_multiplier: CardinalityEstimate
    final: CardinalityEstimate


def spec_cardinality_plan(spec: "Spec") -> SpecCardinalityPlan:
    """Build the full confidence-tagged cardinality plan for `spec` (STEP 9).

    Mirrors `estimate_core_combos`'s mandatory-product walk (same slot filter,
    same Cartes operand sizing) but keeps every stage's confidence tier and
    formula instead of returning one bare integer — see `SpecCardinalityPlan`."""
    sheet_n = {s.sheet: len(s.values) for s in spec.slots}
    raw_values = dict(sheet_n)

    per_slot: "dict[str, CardinalityEstimate]" = {}
    mandatory_estimates: list[CardinalityEstimate] = []
    for s in spec.slots:
        ops = verb_sheet_operands(s.verb)
        other = sheet_n.get(ops[0], 0) if ops else 0
        est = verb_cardinality(s.verb, len(s.values), other)
        per_slot[s.sheet] = est
        if not any(f in s.flags for f in ("FW_Exclude", "FW_Heading", "FW_Optional")):
            mandatory_estimates.append(est)

    # seq_extra brace joiners (STEP 9 fix): a brace FW_(...) replaces its excluded
    # operands' contribution to fw_final with the JOINED result-table row count —
    # an UNKNOWN second-order effect that per-slot estimates can't see. Folding it
    # into the mandatory product is what stops a valid brace spec from collapsing
    # to a bogus "EXACT 1" (empty mandatory product) when both operands are excluded.
    brace_estimates = [brace_cardinality(c) for cells in spec.seq_extra for c in cells
                       if _BRACE_RE.fullmatch(_first_line(c))]
    mandatory_estimates += brace_estimates

    mandatory = _combine_product(
        mandatory_estimates,
        formula="Π verb_cardinality(slot) over mandatory (non-Exclude/Heading/Optional) slots"
                + (" × brace joiner row-count(s) (seq_extra)" if brace_estimates else ""),
        reasons=("matches estimate_core_combos's slot filter and Cartes operand sizing "
                 "(declared sheet size, not post-verb result-table size)",)
                + (("seq_extra brace joiner(s) contribute their joined-result-table row count, "
                    "which is UNKNOWN until Core runs — it cannot be folded in as an exact factor",)
                   if brace_estimates else ()))

    if not spec.constraints:
        post_sieve = CardinalityEstimate(
            mode=mandatory.mode, value=mandatory.value, lower=mandatory.lower, upper=mandatory.upper,
            formula="= mandatory (no constraints declared -> sieve is a no-op)")
    else:
        upper = mandatory.upper if mandatory.upper is not None else mandatory.value
        post_sieve = CardinalityEstimate.bounded(
            lower=0, upper=upper if upper is not None else 0,
            value=upper, formula="[0, mandatory] (sieve predicate selectivity not evaluated)",
            reasons=(f"{len(spec.constraints)} constraint(s) declared; the sieve can remove "
                     f"anywhere from none to all mandatory rows depending on runtime data — "
                     f"only the conservative range is provable without running it",))

    optional_multiplier = optional_multiplier_cardinality(spec)
    final = _combine_product(
        (post_sieve, optional_multiplier),
        formula="post_sieve × optional_multiplier",
        reasons=("each FW_Optional slot independently multiplies the post-sieve candidate space",))

    return SpecCardinalityPlan(raw_values=raw_values, per_slot=per_slot, mandatory=mandatory,
                               post_sieve=post_sieve, optional_multiplier=optional_multiplier, final=final)


def cardinality_estimate_to_dict(est: CardinalityEstimate) -> dict:
    """JSON-able projection of a :class:`CardinalityEstimate` (STEP 9 action 5:
    'Formula доступна в JSON и human output'). Enum -> its string value;
    tuples -> lists, so `json.dumps` needs no custom encoder."""
    return {"mode": est.mode.value, "value": est.value, "lower": est.lower, "upper": est.upper,
            "formula": est.formula, "reasons": list(est.reasons), "assumptions": list(est.assumptions)}


def cardinality_plan_to_dict(plan: SpecCardinalityPlan) -> dict:
    """JSON-able projection of a full :class:`SpecCardinalityPlan`."""
    return {
        "raw_values": dict(plan.raw_values),
        "per_slot": {sheet: cardinality_estimate_to_dict(e) for sheet, e in plan.per_slot.items()},
        "mandatory": cardinality_estimate_to_dict(plan.mandatory),
        "post_sieve": cardinality_estimate_to_dict(plan.post_sieve),
        "optional_multiplier": cardinality_estimate_to_dict(plan.optional_multiplier),
        "final": cardinality_estimate_to_dict(plan.final),
    }


def _format_estimate(label: str, est: CardinalityEstimate) -> str:
    if est.mode == CardinalityMode.EXACT:
        body = f"{est.value}  (EXACT — {est.formula})"
    elif est.mode == CardinalityMode.UNKNOWN:
        body = f"UNKNOWN  ({est.formula})"
    else:
        body = f"{est.value}  ({est.mode.value} [{est.lower}, {est.upper}] — {est.formula})"
    lines = [f"{label}: {body}"]
    for r in est.reasons:
        lines.append(f"    reason: {r}")
    return "\n".join(lines)


def format_cardinality_plan(plan: SpecCardinalityPlan) -> str:
    """Human-readable rendering of a :class:`SpecCardinalityPlan` (STEP 9 action 5:
    'Formula доступна в JSON и human output'). Mirrors `cardinality_plan_to_dict`'s
    stages 1:1 so the two views never drift apart."""
    lines = ["raw values (declared per-slot sizes):"]
    for sheet, n in plan.raw_values.items():
        lines.append(f"  {sheet}: {n}")
    lines.append("per-slot Core rows:")
    for sheet, est in plan.per_slot.items():
        lines.append("  " + _format_estimate(sheet, est).replace("\n", "\n  "))
    lines.append(_format_estimate("mandatory Core product", plan.mandatory))
    lines.append(_format_estimate("estimated post-sieve", plan.post_sieve))
    lines.append(_format_estimate("optional multiplier", plan.optional_multiplier))
    lines.append(_format_estimate("final candidate count", plan.final))
    return "\n".join(lines)


@dataclass(frozen=True)
class Slot:
    sheet: str            # data-sheet name (must NOT start with FW_)
    key: str              # token key used in the assembled cell (' key=value')
    values: list[str]     # row 0 = baseline
    verb: str = DEFAULT_VERB   # Core combination rule for this slot (any is_core_verb token)
    flags: tuple = ()          # extra FW_Seq flags (FW_Optional/FW_Exclude/FW_Heading/…)
    raw: bool = False          # emit values VERBATIM (no ' key=' prefix) — for CODE pieces that
                               # must concatenate into a runnable program; also skips ${} expansion
    prefix: str = ""           # FW_SheetNames prefix wrapper applied Core-side (e.g. a stmt opener)
    ending: str = ""           # FW_SheetNames ending wrapper
    separator: str = ""        # FW_Separator(<sheet>): weave that sheet's FIRST value between elements
                               # (a modifier on this slot's verb; cardinality-preserving)
    group_replace: tuple = ()  # FW_Group: ((pattern, replacement), ...) FW_ReplaceRE rewrites applied to
                               # this slot's produced rows (2nd-order: re-combine/rewrite; see ZEN doc)
    alias: str = ""            # STEP 36: the domain-level authoring alias this slot was compiled FROM
                               # (e.g. "choose_k(2)"); "" = authored with a raw verb. Provenance only —
                               # `verb`/`flags` above are already the compiled (advanced) representation.


@dataclass(frozen=True)
class Goal:
    key: str
    direction: str        # min | max | target


@dataclass(frozen=True)
class CustomVar:
    code: int
    msg: str


class CardinalityMode(str, Enum):
    """Confidence tier for a row-count prediction (STEP 9).

    EXACT     — closed-form formula over known operands (e.g. C(n,k), n!).
    BOUNDED   — true count is unknown but provably within [lower, upper].
    ESTIMATED — best-effort point guess; neither exact nor provably bounded.
    UNKNOWN   — no statically sound number at all; only `reasons` are filled.
    """
    EXACT = "EXACT"
    BOUNDED = "BOUNDED"
    ESTIMATED = "ESTIMATED"
    UNKNOWN = "UNKNOWN"


@dataclass(frozen=True)
class CardinalityEstimate:
    """A row-count prediction that is honest about how sure it is (STEP 9).

    Replaces "one approximate integer" with mode + bounds + the formula and
    reasoning behind it, so a rich verb (FW_Cartes/FW_Group/braces) is never
    silently presented with the same confidence as an exact formula like
    C(n,k). `value` is the single best-effort point estimate (may be None for
    UNKNOWN); `lower`/`upper` bound the true count whenever that is provable.
    """
    mode: CardinalityMode
    value: "int | None" = None
    lower: "int | None" = None
    upper: "int | None" = None
    formula: str = ""
    reasons: tuple[str, ...] = ()
    assumptions: tuple[str, ...] = ()

    @staticmethod
    def exact(value: int, *, formula: str, assumptions: "Sequence[str]" = ()) -> "CardinalityEstimate":
        return CardinalityEstimate(mode=CardinalityMode.EXACT, value=value, lower=value, upper=value,
                                   formula=formula, assumptions=tuple(assumptions))

    @staticmethod
    def bounded(*, lower: int, upper: int, formula: str,
                reasons: "Sequence[str]" = (), assumptions: "Sequence[str]" = (),
                value: "int | None" = None) -> "CardinalityEstimate":
        return CardinalityEstimate(mode=CardinalityMode.BOUNDED, value=value if value is not None else upper,
                                   lower=lower, upper=upper, formula=formula,
                                   reasons=tuple(reasons), assumptions=tuple(assumptions))

    @staticmethod
    def unknown(*, formula: str = "", reasons: "Sequence[str]" = (),
                assumptions: "Sequence[str]" = ()) -> "CardinalityEstimate":
        return CardinalityEstimate(mode=CardinalityMode.UNKNOWN, formula=formula,
                                   reasons=tuple(reasons), assumptions=tuple(assumptions))


@dataclass
class Spec:
    name: str                       # file stem
    title: str
    slots: list[Slot]
    goals: list[Goal] = field(default_factory=list)
    custom_vars: list[CustomVar] = field(default_factory=list)
    args: list[str] = field(default_factory=list)
    note: str = ""
    runme: str = ""
    # Advanced/raw extra FW_Seq rows (each a list of cell strings) for constructs
    # the per-slot model can't express directly — brace joiners FW_(…), multi-verb
    # rows, nested FW_(…FW_()…). Appended verbatim after the slot rows; their sheet
    # references are still validated against the declared data sheets.
    seq_extra: list = field(default_factory=list)
    # Constraint sidecar (the "allowed/forbidden bonds" layer; enforced by constraints/sieve.py
    # between Core and Reader, NOT by the Core DSL). params: {sheet:{value:{attr:val}}};
    # constraints: list of {id,sheets,(pairs|when),gate,polarity,desc}. See constraints/sidecar_schema.md.
    params: dict = field(default_factory=dict)
    constraints: list = field(default_factory=list)
    # Ordinal `orders` for the sieve's ordinal leaves (ge/le/…, geSheet/…): {sheet: [v0,v1,…] |
    # "numeric" | "date"}. Flows verbatim into the sidecar; see constraints/sidecar_schema.md.
    orders: dict = field(default_factory=dict)
    # STEP 8: explicit authoring-contract version. Specs predating this field
    # carry no `spec_version` key and are interpreted as "legacy" (today's
    # format, unchanged); `"1"` is the same internal model with the contract
    # made explicit — see bundle-spec-v1.schema.json.
    spec_version: str = "legacy"

    @property
    def baseline(self) -> tuple[str, ...]:
        return tuple(s.values[0] for s in self.slots)

    @property
    def combos(self) -> int:
        n = 1
        for s in self.slots:
            n *= len(s.values)
        return n

    def goals_property(self) -> str:
        return " ; ".join(f"{'maximize' if g.direction=='max' else 'minimize' if g.direction=='min' else 'target'} {g.key}"
                          for g in self.goals)


# STEP 8: top-level spec keys recognized by parse_spec/bundle-spec-v1.schema.json.
# Used only to flag authoring typos (strict mode rejects, compatibility mode warns);
# the parser itself stays driven by `.get(...)` defaults, so this list is additive
# documentation, not a new validation path for already-accepted specs.
SPEC_V1_KNOWN_KEYS = frozenset({
    "spec_version", "title", "note", "args", "runme", "slots",
    "goals", "custom_vars", "params", "constraints", "seq_extra", "orders",
})

# Per-slot keys recognized inside `slots[]` entries — same typo-detection role as
# SPEC_V1_KNOWN_KEYS, one level down (e.g. catches `flgas` instead of `flags`).
SPEC_V1_SLOT_KNOWN_KEYS = frozenset({
    "sheet", "key", "values", "verb", "flags", "raw",
    "prefix", "ending", "separator", "group_replace",
    "alias",   # STEP 36: domain-level authoring alias (compiles to verb+flags)
})

# spec_version values this loader understands. Anything else fails closed —
# an unrecognized contract version is a hard authoring error, not a style nit,
# so this check applies regardless of strict/compatibility mode.
SUPPORTED_SPEC_VERSIONS = frozenset({"legacy", "1"})

# Nested-object key allowlists — same typo-detection role as SPEC_V1_*_KNOWN_KEYS,
# applied to dict-shaped entries inside goals[]/custom_vars[]/constraints[].
# (params[] rows intentionally carry open-ended <attr>=... pairs and are NOT
# checked here — that openness is part of their contract, see emit_sidecar.)
SPEC_V1_GOAL_KNOWN_KEYS = frozenset({"key", "dir", "direction"})
SPEC_V1_CUSTOM_VAR_KNOWN_KEYS = frozenset({"code", "msg"})
SPEC_V1_CONSTRAINT_KNOWN_KEYS = frozenset({"id", "sheets", "pairs", "sets", "when", "gate", "polarity", "desc",
                                           "condition", "mapping", "assert"})


def _condition_sheets(cond) -> list:
    """Sheets a constraint `condition` AST references (mirrors constraints/sieve._condition_sheets);
    used to validate that a conditional bond's context sheets are declared."""
    if not cond:
        return []
    out: list = []
    for key in ("all", "any"):
        for c in cond.get(key, []) or []:
            out += _condition_sheets(c)
    if cond.get("not"):
        out += _condition_sheets(cond["not"])
    for key in ("sheet",) + tuple(_COND_SHEET_OPS):
        if cond.get(key):
            out.append(cond[key])
    return out


# Mirror of constraints/sieve.py's condition op-sets — KEEP IN SYNC (the sieve is the source of truth).
_COND_NODE_KEYS = {"all", "any", "not"}
_COND_SET_OPS = {"in", "nin", "hasAny"}
_COND_COUNT_OPS = {"count", "countGe", "countLe"}
_COND_SHEET_OPS = {"eqSheet", "neSheet", "geSheet", "gtSheet", "leSheet", "ltSheet", "subOf", "supOf"}
_COND_ORD_CONST_OPS = {"ge", "gt", "le", "lt"}
_COND_LEAF_OPS = (
    {"eq", "ne", "has", "hasnt", "in", "nin", "hasAny", "present"}
    | _COND_SET_OPS | _COND_COUNT_OPS | _COND_SHEET_OPS | _COND_ORD_CONST_OPS
)
_COND_ALLOWED_KEYS = _COND_NODE_KEYS | _COND_LEAF_OPS | {"sheet"}


def _condition_error(cond, path: str = "condition") -> str | None:
    if cond is None or cond == {}:
        return None
    if not isinstance(cond, dict):
        return f"{path} must be an object"
    unknown = sorted(set(cond) - _COND_ALLOWED_KEYS)
    if unknown:
        return f"{path} has unknown key(s) {unknown}"
    node_keys = [k for k in _COND_NODE_KEYS if k in cond]
    leaf_ops = [k for k in _COND_LEAF_OPS if k in cond]
    if node_keys:
        if len(node_keys) != 1 or leaf_ops or "sheet" in cond:
            return f"{path} must be either one boolean node or one sheet leaf"
        key = node_keys[0]
        if key in ("all", "any"):
            children = cond.get(key)
            if not isinstance(children, list) or not children:
                return f"{path}.{key} must be a non-empty list"
            for i, child in enumerate(children):
                err = _condition_error(child, f"{path}.{key}[{i}]")
                if err:
                    return err
        else:
            err = _condition_error(cond.get("not"), f"{path}.not")
            if err:
                return err
        return None
    if "sheet" not in cond:
        return f"{path} leaf needs 'sheet'"
    if len(leaf_ops) != 1:
        return f"{path} leaf needs exactly one operator, got {leaf_ops or 'none'}"
    op = leaf_ops[0]
    val = cond.get(op)
    if op in _COND_SET_OPS and (isinstance(val, (str, bytes)) or not isinstance(val, (list, tuple, set))):
        return f"{path}.{op} must be a list of values"
    if op in _COND_COUNT_OPS and not isinstance(val, int):
        return f"{path}.{op} must be an integer"
    if op in _COND_SHEET_OPS and not (isinstance(val, str) and val):
        return f"{path}.{op} must name another sheet"
    if op == "present" and not isinstance(val, bool):
        return f"{path}.present must be true or false"
    if op in _COND_ORD_CONST_OPS and (val is None or isinstance(val, (list, tuple, set, dict))):
        return f"{path}.{op} must be a scalar constant (rank via `orders`)"
    return None


def _mapping_error(m, path: str = "mapping") -> str | None:
    if m is None:
        return None
    if not isinstance(m, dict):
        return f"{path} must be an object"
    unknown = sorted(set(m) - {"source", "target", "allow"})
    if unknown:
        return f"{path} has unknown key(s) {unknown}"
    if not m.get("source") or not m.get("target"):
        return f"{path} needs non-empty source and target"
    allow = m.get("allow", {})
    if not isinstance(allow, dict):
        return f"{path}.allow must be an object"
    for key, vals in allow.items():
        if isinstance(vals, (str, bytes)) or not isinstance(vals, (list, tuple, set)):
            return f"{path}.allow[{key!r}] must be a list of values"
    return None


def _check_unknown_keys(d: dict, known: frozenset, *, where: str, name: str, strict: bool) -> None:
    unknown = sorted(set(d) - known)
    if not unknown:
        return
    if strict:
        raise ValueError(f"spec '{name}' {where} has unknown field(s) {unknown} (strict mode)")
    print(f"  ⚠ spec '{name}': {where} unknown field(s) {unknown} ignored (compatibility mode)")


def _default_runme(label: str) -> str:
    return ('class RunMeFirstOnce { public static String FW_ARGS; '
            'public static void main(String[] args){ FW_ARGS = ""; '
            'System.out.println("%s search space"); } }' % label)


def _load_raw(path: Path) -> dict:
    ext = path.suffix.lower()
    text = path.read_text(encoding="utf-8")
    if ext == ".toml":
        import tomllib
        return tomllib.loads(text)
    if ext == ".json":
        return json.loads(text)
    if ext in (".yaml", ".yml"):
        import yaml
        return yaml.safe_load(text)
    raise ValueError(f"Unsupported spec format: {path.name} (use .toml/.json/.yaml)")


def parse_spec(raw: dict, name: str, *, strict: bool = False) -> Spec:
    """Turn a raw dict into a Spec, filling in everything the author left out
    (the 'smart picks the variants' part). The ONLY required field is `slots`.

    STEP 8: `spec_version` is optional and defaults to `"legacy"` (today's
    format, normalized to the same internal model as `"1"` — see
    bundle-spec-v1.schema.json). An unrecognized `spec_version` (e.g. `"999"`)
    always fails closed, in either mode — that's a hard contract mismatch, not
    a style nit. Unknown keys — top-level, per-slot, and inside dict-shaped
    `goals[]`/`custom_vars[]`/`constraints[]` entries — are a no-op for the
    existing legacy interpretation; `strict=True` rejects them outright (typo
    detection, e.g. a slot's `flgas` instead of `flags`, or a goal's
    `direktion` instead of `direction`), otherwise they only print a
    compatibility warning. (`params[]` rows intentionally allow open-ended
    `<attr>=...` pairs and are exempt — see emit_sidecar.)
    """
    spec_version = str(raw.get("spec_version", "legacy"))
    if spec_version not in SUPPORTED_SPEC_VERSIONS:
        raise ValueError(f"spec '{name}' has unsupported spec_version {spec_version!r} "
                         f"(supported: {sorted(SUPPORTED_SPEC_VERSIONS)})")
    _check_unknown_keys(raw, SPEC_V1_KNOWN_KEYS, where="top level", name=name, strict=strict)
    slots: list[Slot] = []
    for s in raw.get("slots", []):
        _check_unknown_keys(s, SPEC_V1_SLOT_KNOWN_KEYS,
                            where=f"slot '{s.get('sheet', '?')}'", name=name, strict=strict)
        sheet = s["sheet"]
        if sheet.startswith("FW_"):
            raise ValueError(f"slot sheet '{sheet}' must not start with FW_ (it would be read as a control sheet)")
        is_raw = bool(s.get("raw", False))
        str_vals = [str(v) for v in s["values"]]
        # raw (code) slots are taken VERBATIM — no ${a|b} expansion (code may contain ${...})
        exp_vals = str_vals if is_raw else [ev for v in str_vals for ev in expand_options(v)]
        # STEP 36: a slot is authored with EITHER a domain `alias` OR a raw `verb`,
        # never both — a conflicting alias+verb is rejected. The alias compiles to
        # the same advanced verb (+flags) an expert would write; `alias_token`
        # records the provenance for the normalized artifact.
        alias_token = str(s.get("alias", "")).strip()
        alias_flags: tuple[str, ...] = ()
        if alias_token:
            if "verb" in s:
                raise ValueError(f"slot '{sheet}' sets BOTH alias {alias_token!r} and raw verb "
                                 f"{str(s['verb'])!r} — choose one (an alias already compiles to a "
                                 f"verb; it cannot also carry an explicit verb)")
            try:
                verb, alias_flags = compile_alias(alias_token)
            except ValueError as e:
                raise ValueError(f"slot '{sheet}': {e}")
        else:
            verb = str(s.get("verb", DEFAULT_VERB)).strip()
        if not is_core_verb(verb):
            raise ValueError(f"slot '{sheet}' verb {verb!r} is not a recognized Core FW_Seq verb "
                             f"(e.g. FW_Combi(2), FW_CombiR(2), FW_Permut, FW_PermutR(2), "
                             f"FW_Subsets, FW_Subsets_EXACT(2), FW_Cartes(OTHER), FW_Group)")
        flags = tuple(str(f).strip() for f in s.get("flags", []))
        # alias-contributed flags (e.g. optional_action -> FW_Optional) merge in, de-duped,
        # appended after any author-listed flags so the compiled order is deterministic.
        flags = flags + tuple(f for f in alias_flags if f not in flags)
        bad = [f for f in flags if not is_core_flag(f)]
        if bad:
            raise ValueError(f"slot '{sheet}' has unrecognized flag(s) {bad}")
        separator = str(s.get("separator", "")).strip()
        gr_raw = s.get("group_replace", [])
        for e in gr_raw:
            if len(e) != 2:
                raise ValueError(f"slot '{sheet}' group_replace entries must be [pattern, replacement], got {e!r}")
        group_replace = tuple((str(e[0]), str(e[1])) for e in gr_raw)
        slots.append(Slot(sheet=sheet,
                          key=s.get("key", sheet.lower()),
                          values=exp_vals, verb=verb, flags=flags,
                          raw=is_raw, prefix=str(s.get("prefix", "")), ending=str(s.get("ending", "")),
                          separator=separator, group_replace=group_replace, alias=alias_token))
    if not slots:
        raise ValueError(f"spec '{name}' has no slots")

    # seq_extra: raw advanced rows (brace joiners / multi-verb / nested). Each row
    # is a list of cell strings; validate it carries a recognized verb/joiner and
    # that any referenced sheets are declared as slots.
    sheet_names = {s.sheet for s in slots}
    flags_by_sheet = {s.sheet: s.flags for s in slots}
    for s in slots:                                     # FW_Separator operand must be a declared sheet
        if s.separator and s.separator not in sheet_names:
            raise ValueError(f"slot '{s.sheet}' separator '{s.separator}' is not a declared sheet")
    seq_extra: list[list[str]] = []
    for row in raw.get("seq_extra", []):
        cells = [str(c) for c in row if str(c).strip() != ""]
        if not cells:
            continue
        if not any(is_core_verb(c) for c in cells):
            raise ValueError(f"spec '{name}' seq_extra row {cells} has no recognized Core verb/joiner")
        for c in cells:
            for op in verb_sheet_operands(c):
                if op not in sheet_names:
                    raise ValueError(f"spec '{name}' seq_extra references sheet '{op}' that is not a declared slot")
            # Core contract: a brace FW_(...) JOINS the RESULT TABLES of its two
            # operands (positions 2 & 4) — the combination-results that the prior
            # FW_Seq rows produced for those sheets (fw2_<k>/fw_<k>), not the raw
            # sheets — via the exclude map. So those operand slots MUST be set aside
            # with FW_Exclude (or FW_Heading); otherwise they ALSO remain independent
            # axes in fw_final and the join double-counts / empties them. (Verified in
            # SeqParser.parseExcludedKeys + BraceOperationHandler.resolveSourceTableForInner
            # + cleanupExcludedTables.)
            first = _first_line(c)
            if _BRACE_RE.fullmatch(first):
                inner = first[first.index("(") + 1:first.rindex(")")].split(",")
                for i in (2, 4):
                    op = inner[i] if i < len(inner) else ""
                    if op in {"FW_()", "FW_()G"}:
                        continue
                    if op and not ({"FW_Exclude", "FW_Heading"} & set(flags_by_sheet.get(op, ()))):
                        raise ValueError(
                            f"spec '{name}' brace {first!r} joins operand '{op}', but that slot is "
                            f"not FW_Exclude'd. Brace operands must carry flags=['FW_Exclude'] (set "
                            f"aside from the cartesian) so they appear only via the join — add it, or "
                            f"the Core double-counts the operand. (Core: SeqParser/BraceOperationHandler.)")
        seq_extra.append(cells)

    # params: [{sheet, value, <attr>=...}] -> {sheet: {value: {attrs}}} (for predicate constraints)
    params: dict = {}
    for row in raw.get("params", []):
        sh, val = row.get("sheet"), str(row.get("value"))
        if sh not in sheet_names:
            raise ValueError(f"spec '{name}' params reference undeclared sheet '{sh}'")
        params.setdefault(sh, {})[val] = {k: v for k, v in row.items() if k not in ("sheet", "value")}

    # constraints: allowed/forbidden bonds (enforced by constraints/sieve.py between Core & Reader)
    constraints: list = []
    for c in raw.get("constraints", []):
        _check_unknown_keys(c, SPEC_V1_CONSTRAINT_KNOWN_KEYS,
                            where=f"constraint '{c.get('id', '?')}'", name=name, strict=strict)
        cond_err = _condition_error(c.get("condition"))
        if cond_err:
            raise ValueError(f"spec '{name}' constraint {c.get('id')!r} has invalid condition: {cond_err}")
        asrt_err = _condition_error(c.get("assert"), "assert") if c.get("assert") is not None else None
        if asrt_err:
            raise ValueError(f"spec '{name}' constraint {c.get('id')!r} has invalid assert: {asrt_err}")
        mp_raw = c.get("mapping")
        map_err = _mapping_error(mp_raw)
        if map_err:
            raise ValueError(f"spec '{name}' constraint {c.get('id')!r} has invalid mapping: {map_err}")
        mp = mp_raw or {}
        cs = (c.get("sheets") or list(c.get("sets", {}).keys())
              or ([mp["source"], mp["target"]] if mp else [])
              or sorted({k for e in c.get("pairs", []) for k in e}))
        # validate EVERY referenced sheet is declared: the target sheets, the mapping's
        # source/target, the `assert` body sheets, and the `condition` context sheets (a conditional
        # bond touches all of them).
        ref = set(cs)
        if mp:
            ref |= {mp.get("source"), mp.get("target")}
        ref |= set(_condition_sheets(c.get("condition")))
        ref |= set(_condition_sheets(c.get("assert")))
        for sh in ref:
            if sh and sh not in sheet_names:
                raise ValueError(f"spec '{name}' constraint {c.get('id')!r} references undeclared sheet '{sh}'")
        # `assert` (a condition-AST as the bond body) is a valid standalone bond family.
        if not any(k in c for k in ("when", "pairs", "sets", "mapping", "assert")):
            raise ValueError(f"spec '{name}' constraint {c.get('id')!r} needs 'pairs', 'sets', 'when', 'mapping', or 'assert'")
        constraints.append(c)

    # orders: {sheet: [v0,v1,…] | "numeric" | "date"} — ordinal ranks for the sieve's ordinal leaves.
    orders: dict = {}
    for sh, spec_ord in (raw.get("orders", {}) or {}).items():
        if sh not in sheet_names:
            raise ValueError(f"spec '{name}' orders reference undeclared sheet '{sh}'")
        if not (isinstance(spec_ord, (list, tuple)) or spec_ord in ("numeric", "date")):
            raise ValueError(f"spec '{name}' orders[{sh!r}] must be a list, 'numeric', or 'date'")
        orders[sh] = list(spec_ord) if isinstance(spec_ord, (list, tuple)) else spec_ord

    # goals: accept ["f1","latency_ms"] (auto-direction) or [{key,dir}].
    goals: list[Goal] = []
    for g in raw.get("goals", []):
        if isinstance(g, str):
            goals.append(Goal(g, infer_direction(g)))
        else:
            _check_unknown_keys(g, SPEC_V1_GOAL_KNOWN_KEYS,
                                where=f"goal '{g.get('key', '?')}'", name=name, strict=strict)
            goals.append(Goal(g["key"], g.get("dir") or g.get("direction") or infer_direction(g["key"])))

    # custom_vars: explicit list, else auto-derive generic codes from 2.
    cvs: list[CustomVar] = []
    for c in raw.get("custom_vars", []):
        _check_unknown_keys(c, SPEC_V1_CUSTOM_VAR_KNOWN_KEYS,
                            where=f"custom_var (code={c.get('code', '?')})", name=name, strict=strict)
        cvs.append(CustomVar(int(c["code"]), str(c["msg"])))
    if not cvs:
        cvs = [CustomVar(2, "candidate failed verdict (nonzero FW_CUSTOM_VAR)")]

    return Spec(name=name,
                title=raw.get("title", name),
                slots=slots,
                goals=goals,
                custom_vars=cvs,
                args=[str(a) for a in raw.get("args", [])] or ["noargs"],
                note=raw.get("note", ""),
                runme=raw.get("runme") or _default_runme(name),
                seq_extra=seq_extra, params=params, constraints=constraints, orders=orders,
                spec_version=spec_version)


def emit_sidecar(spec: Spec, path: str | Path) -> dict:
    """Write the constraint sidecar (params + constraints) as JSON for constraints/sieve.py.
    The sieve runs this between Core and Reader; see constraints/sidecar_schema.md."""
    import json
    sc = {"version": 1, "params": spec.params, "constraints": spec.constraints}
    if getattr(spec, "orders", None):
        sc["orders"] = spec.orders
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(sc, indent=2, ensure_ascii=False), encoding="utf-8")
    return sc


def normalized_spec_dict(spec: Spec) -> dict:
    """STEP 36: the COMPILED (advanced) representation of `spec` as a raw spec dict.

    Any domain-level alias (``choose_one``/``choose_k``/``permute``/
    ``feature_subset``/``optional_action``) has already been lowered to its
    expert ``verb`` (+flags) by `parse_spec`; this dict carries ONLY raw verbs
    and round-trips through `parse_spec` to an equivalent Spec. It is what gets
    saved in run artifacts so the advanced form an alias expands to is
    inspectable (and a reviewer can diff alias-vs-expert specs)."""
    slots: list[dict] = []
    for s in spec.slots:
        d: dict = {"sheet": s.sheet, "key": s.key, "values": list(s.values), "verb": s.verb}
        if s.flags:
            d["flags"] = list(s.flags)
        if s.raw:
            d["raw"] = True
        if s.prefix:
            d["prefix"] = s.prefix
        if s.ending:
            d["ending"] = s.ending
        if s.separator:
            d["separator"] = s.separator
        if s.group_replace:
            d["group_replace"] = [list(e) for e in s.group_replace]
        slots.append(d)
    # params {sheet:{value:{attrs}}} -> the list-of-rows authoring form
    params_rows = [({"sheet": sh, "value": val} | dict(attrs))
                   for sh, vals in spec.params.items() for val, attrs in vals.items()]
    out: dict = {
        "spec_version": spec.spec_version,
        "title": spec.title,
        "runme": spec.runme,   # preserve verbatim so the dict round-trips faithfully
        "slots": slots,
        "goals": [{"key": g.key, "dir": g.direction} for g in spec.goals],
        "custom_vars": [{"code": c.code, "msg": c.msg} for c in spec.custom_vars],
    }
    if spec.note:
        out["note"] = spec.note
    if spec.args and spec.args != ["noargs"]:
        out["args"] = list(spec.args)
    if spec.seq_extra:
        out["seq_extra"] = [list(r) for r in spec.seq_extra]
    if params_rows:
        out["params"] = params_rows
    if spec.constraints:
        out["constraints"] = list(spec.constraints)
    return out


def emit_normalized_spec(spec: Spec, path: str | Path) -> dict:
    """Write the normalized advanced spec (see `normalized_spec_dict`) to JSON for
    run artifacts. The file wraps the round-trippable ``normalized`` dict plus an
    ``alias_provenance`` list mapping each aliased slot to the verb/flags it
    compiled to (for a quick inspectable audit). Returns the written object."""
    import json
    norm = normalized_spec_dict(spec)
    provenance = [{"sheet": s.sheet, "alias": s.alias, "verb": s.verb, "flags": list(s.flags)}
                  for s in spec.slots if s.alias]
    obj = {"spec": spec.name, "normalized": norm, "alias_provenance": provenance}
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(obj, indent=2, ensure_ascii=False), encoding="utf-8")
    return obj


def load_spec(path: str | Path, *, strict: bool = False) -> Spec:
    path = Path(path)
    return parse_spec(_load_raw(path), path.stem, strict=strict)


_SPEC_FILE_SUFFIXES = {".toml", ".json", ".yaml", ".yml"}
_IGNORED_SPEC_SUFFIXES = {".jsonl"}
_BUNDLE_ARTIFACT_JSON_NAMES = {
    "plan.json",
    "run.json",
    "stage.json",
    "handoff.json",
    "state.json",
    "metrics.json",
    "result_card.json",
    "scenario_evidence.json",
    "release_gate.json",
    "verification_report.json",
    "executor-summary.json",
    "zen_mapping.json",
}


def _is_bundle_artifact_json(path: Path) -> bool:
    """True if *path* is one of the Bundle's OWN JSON outputs, not an authored spec.

    `bundle plan <dir>` writes plan.json into the spec dir by default, and newer
    self-posed scenario tooling may leave result cards / evidence beside specs. The
    dir loader accepts .json, so these files must be ignored before parse_spec sees
    them. Constraint-editor sidecars are also JSON and are often downloaded beside
    the spec while authoring; those must not be parsed as specs either. Detection is
    intentionally conservative: known artifact names, meta/ JSON convention,
    *.graph.json / *.sidecar*.json files, stamped top-level bundle./analyzer.
    schemas, and the exact v1 sidecar shape with no slots.
    """
    if path.suffix.lower() != ".json":
        return False
    name = path.name
    if (path.parent.name == "meta"
            or name in _BUNDLE_ARTIFACT_JSON_NAMES
            or name.endswith(".graph.json")
            or name == "sidecar.json"
            or ".sidecar" in name):
        return True
    try:
        with path.open(encoding="utf-8") as f:
            doc = json.load(f)
    except (OSError, ValueError):
        return False
    if not isinstance(doc, dict):
        return False
    if ("slots" not in doc
            and doc.get("version") == 1
            and isinstance(doc.get("constraints"), list)
            and set(doc).issubset({"version", "params", "constraints"})):
        return True
    schema = doc.get("schema")
    return isinstance(schema, str) and (schema.startswith("bundle.") or schema.startswith("analyzer."))


def _is_spec_dir_entry(path: Path) -> bool:
    if not path.is_file():
        return False
    suffix = path.suffix.lower()
    if suffix in _IGNORED_SPEC_SUFFIXES:
        return False
    if suffix not in _SPEC_FILE_SUFFIXES:
        return False
    return not _is_bundle_artifact_json(path)


def load_specs_dir(d: str | Path, *, strict: bool = False) -> list[Spec]:
    d = Path(d)
    specs = [load_spec(p, strict=strict) for p in sorted(d.iterdir()) if _is_spec_dir_entry(p)]
    if not specs:
        raise ValueError(f"no spec files (.toml/.json/.yaml) found in {d}")
    return specs


# ===========================================================================
# 3. Generation engine  (itertools, from v25 combiner_core — structured)
# ===========================================================================

_OPT_RE = re.compile(r"(?<!\\)\$\{([^}]+)\}")   # ${a|b|c} option group


def expand_options(value: str) -> list[str]:
    """Expand a value containing ${a|b|c} option groups into all variants
    (v25 feature). 'x=${1|2}' -> ['x=1','x=2']. No groups -> [value]."""
    parts = _OPT_RE.split(value)
    if len(parts) == 1:
        return [value]
    choices = [[p] if i % 2 == 0 else p.split("|") for i, p in enumerate(parts)]
    return ["".join(combo) for combo in itertools.product(*choices)]


def cartesian(spec: Spec) -> Iterator[tuple[str, ...]]:
    """Full Cartesian product of the slots (values were ${a|b}-expanded at load)."""
    return itertools.product(*[s.values for s in spec.slots])


def assemble(spec: Spec, combo: Sequence[str], prefix: str = "") -> str:
    """One candidate spec line: ' key=value' per slot (leading space = separator;
    never starts with FW_). Optional ' ablation=<label>' prefix."""
    body = "".join(f" {s.key}={v}" for s, v in zip(spec.slots, combo))
    return (f" {prefix}" if prefix else "") + body


# ===========================================================================
# 4. OPT-IN N-wise reduction  (adapted from v25 combiner_reduction)
# ---------------------------------------------------------------------------
# Operates on STRUCTURED combos (list of value-tuples), so position = slot index
# and there is no fragile string re-tokenisation (an improvement over v25, which
# reduced text files). Default is NO reduction (exhaustive); reduction is always
# an explicit, user-chosen step.
# ===========================================================================

def _ntuples(combo: Sequence[str], n: int):
    return itertools.combinations(enumerate(combo), n)


def nwise_greedy(combos: list[tuple[str, ...]], n: int) -> list[tuple[str, ...]]:
    """Streaming greedy N-wise: keep a combo iff it introduces a new
    (slot_index, value) n-tuple. One pass, fast."""
    if n <= 0:
        return list(combos)
    seen: set = set()
    kept: list[tuple[str, ...]] = []
    for c in combos:
        if len(c) < n:
            kept.append(c)
            continue
        tents = list(_ntuples(c, n))
        if any(t not in seen for t in tents):
            seen.update(tents)
            kept.append(c)
    log.info("nwise_greedy n=%d: %d -> %d (covers %d %d-tuples)",
             n, len(combos), len(kept), len(seen), n)
    return kept


def nwise_optimal(combos: list[tuple[str, ...]], n: int) -> list[tuple[str, ...]]:
    """Two-pass greedy minimum-cover N-wise (set cover). Builds the universe of
    all n-tuples, then repeatedly takes the combo covering the most uncovered
    tuples. Smaller suite than greedy, at the cost of holding tuples in RAM."""
    if n <= 0:
        return list(combos)
    tup_sets = []
    universe: set = set()
    short = []
    for i, c in enumerate(combos):
        if len(c) < n:
            short.append(i)
            tup_sets.append(set())
            continue
        ts = set(_ntuples(c, n))
        tup_sets.append(ts)
        universe |= ts
    uncovered = set(universe)
    keep = set(short)
    while uncovered:
        best_i, best_gain = -1, 0
        for i, ts in enumerate(tup_sets):
            if i in keep:
                continue
            g = len(ts & uncovered)
            if g > best_gain:
                best_gain, best_i = g, i
        if best_i < 0 or best_gain == 0:
            break
        keep.add(best_i)
        uncovered -= tup_sets[best_i]
    out = [combos[i] for i in sorted(keep)]
    log.info("nwise_optimal n=%d: %d -> %d (universe %d %d-tuples)",
             n, len(combos), len(out), len(universe), n)
    return out


def reduce_combos(combos: list[tuple[str, ...]], n: int, optimal: bool = False):
    return nwise_optimal(combos, n) if optimal else nwise_greedy(combos, n)


def pick_n_for_budget(spec: Spec, budget: int, optimal: bool = False,
                      hard_cap: int = 2_000_000) -> tuple[int, list[tuple[str, ...]]]:
    """'Self-limit with the exponential wall in mind': choose the MOST thorough
    coverage that fits `budget`. Returns (n, combos) where n=0 means full.
    Refuses to materialise beyond hard_cap (raises) — the unreachable zone."""
    full = spec.combos
    if full <= budget:
        return 0, list(cartesian(spec))
    if full > hard_cap:
        raise ValueError(f"full product {full:,} exceeds hard cap {hard_cap:,}; "
                         f"give an explicit --n for N-wise reduction")
    allc = list(cartesian(spec))
    k = len(spec.slots)
    for n in range(k - 1, 0, -1):           # most-thorough that fits, high->low
        red = reduce_combos(allc, n, optimal)
        if len(red) <= budget:
            return n, red
    return 1, reduce_combos(allc, 1, optimal)


# ===========================================================================
# 5. Workbook builder  (Core-faithful FW_ formats; XLSX + JSON)
# ===========================================================================

def _write_control_sheets(wb: Workbook, spec: Spec, data_sheets: list[str],
                          fw_seq_ws, fw_info: str, clone_path: Optional[str]):
    """Add FW_SheetNames, FW_RunMeFirstOnce, FW_Arguments, FW_CUSTOM_VAR, FW_Info.
    (FW_Seq is built by the caller and passed in for the FW_Info 'dup' mode.)"""
    pe = {s.sheet: (s.prefix, s.ending) for s in spec.slots}   # per-slot FW_SheetNames wrappers
    nm = wb.create_sheet("FW_SheetNames")
    for i, s in enumerate(data_sheets, start=1):
        pre, end = pe.get(s, ("", ""))
        nm.cell(i, 1, s)
        nm.cell(i, 2, pre if pre else "FW_EMPTY_STRING")       # wrappers written VERBATIM (code-safe)
        nm.cell(i, 3, end if end else "FW_EMPTY_STRING")

    wb.create_sheet("FW_RunMeFirstOnce").cell(1, 1, spec.runme)

    aw = wb.create_sheet("FW_Arguments")
    for i, a in enumerate(spec.args, start=1):
        aw.cell(i, 1, a)

    cv = wb.create_sheet("FW_CUSTOM_VAR")
    for i, c in enumerate(spec.custom_vars, start=1):
        cv.cell(i, 1, int(c.code))                         # MUST be int
        cv.cell(i, 2, f"FWCUSTOMVAR={c.code} {c.msg}")

    _add_fw_info(wb, fw_seq_ws, fw_info, clone_path)


def _add_fw_info(wb: Workbook, fw_seq_ws, fw_info: str, clone_path: Optional[str]):
    """FW_Info is inert documentation. Modes: 'dup' (copy FW_Seq, v25 style),
    'clone:<path>' (copy a reference workbook's canonical legend), 'skip'."""
    if fw_info == "skip":
        return
    ws = wb.create_sheet("FW_Info")
    if fw_info == "dup":
        for row in fw_seq_ws.iter_rows(values_only=True):
            ws.append(row)
    elif fw_info == "clone" and clone_path:
        try:
            src = openpyxl.load_workbook(clone_path, read_only=True, data_only=False)
            sn = next((n for n in ("FW_Info", "FW_info") if n in src.sheetnames), None)
            if sn:
                for r in src[sn].iter_rows():
                    for c in r:
                        if c.value not in (None, ""):
                            ws.cell(c.row, c.column, c.value)
            src.close()
        except Exception as e:                              # noqa: BLE001
            log.warning("FW_Info clone from %s failed (%s); leaving empty", clone_path, e)


def _group_directive(replaces) -> str:
    r"""Build the multi-line FW_Group directive cell: 'FW_Group' then one
    FW_ReplaceRE("pattern", "replacement") line per rewrite. Mirrors the proven
    layout in testgen_api.py and the Core regex FW_ReplaceRE\(["](.+?)["],\s+("(.*)")\)."""
    lines = ["FW_Group"]
    for pat, rep in replaces:
        lines.append(f'FW_ReplaceRE("{pat}", "{rep}")')
    return "\n".join(lines)


def _gate_seq_graph(spec: Spec) -> None:
    """STEP 37: build the FW_Seq dependency graph and reject the spec BEFORE any
    workbook is emitted if it has an error-level issue (cycle / missing operand /
    consumed-result ambiguity). Lazy import keeps fwgen free of an fwseq_graph
    import cycle (fwseq_graph imports fwgen)."""
    from fwseq_graph import validate_spec_graph
    validate_spec_graph(spec)


def build_compact(spec: Spec, fw_info: str = "dup",
                  clone_path: Optional[str] = None, autofit: float = 0.0) -> Workbook:
    """Compact workbook: each slot is a FW_Reuse data sheet carrying its chosen
    Core combination verb (default FW_Combi(1) = cartesian leaves), so the CORE
    expands the space. Any `spec.seq_extra` rows (brace joiners / multi-verb /
    nested) are appended verbatim. Small file, native, scales to any product size.

    Layout matches the Core convention 'col B,C = flag slots, col D+ = verbs'; for
    a default FW_Combi(1) slot with no extra flags this is byte-identical to before
    (B=FW_Reuse, D=FW_Combi(1))."""
    _gate_seq_graph(spec)                                   # STEP 37: reject cycle/missing-operand BEFORE the workbook
    wb = Workbook()
    seq = wb.active
    seq.title = "FW_Seq"
    row = 1
    for s in spec.slots:
        seq.cell(row, 1, s.sheet)
        seq.cell(row, 2, "FW_Reuse")                       # no-op for a plain slot; lets it also serve as a brace operand (table kept through join cleanup)
        col = 3
        for f in s.flags:
            if f != "FW_Reuse":
                seq.cell(row, col, f); col += 1            # extra routing flags (Optional/Exclude/…)
        vcol = max(col, 4)
        seq.cell(row, vcol, cell_for_xlsx(s.verb))         # the combination verb (col D+)
        mcol = vcol + 1                                     # modifiers/2nd-order directives follow the verb
        if s.separator:
            seq.cell(row, mcol, f"FW_Separator({s.separator})"); mcol += 1
        if s.group_replace:
            seq.cell(row, mcol, _group_directive(s.group_replace)); mcol += 1   # multi-line, written raw
        row += 1
    for cells in spec.seq_extra:
        for c, val in enumerate(cells, start=1):
            seq.cell(row, c, cell_for_xlsx(val))
        row += 1
    _write_control_sheets(wb, spec, [s.sheet for s in spec.slots], seq, fw_info, clone_path)
    for s in spec.slots:
        ws = wb.create_sheet(s.sheet)
        for r, v in enumerate(s.values, start=1):
            # raw (code) slots: write VERBATIM (cell_for_xlsx's \n/\t→char substitution would
            # corrupt code that contains literal backslash-escapes). non-raw: ' key=value' as before.
            ws.cell(r, 1, v if s.raw else cell_for_xlsx(f" {s.key}={v}"))
    autofit_workbook(wb, autofit)
    return wb


def build_compact_core_json(spec: Spec, fw_info: str = "skip") -> dict:
    r"""Same logical workbook as build_compact(), but emitted as the Core's native
    JSON schedule { "sheets": { name: [[cell,...], ...] } } that JsonScheduleParser
    reads (excel.file=<f>.json + core.input.format=json). Built DIRECTLY from the
    spec — the authoring path never touches openpyxl/POI, so the spec FILE has no
    cell-size limit. (NB: the Core's JsonScheduleParser still materialises into a POI
    XSSFWorkbook downstream, which caps each CELL at 32,767 chars — so pre-split
    oversize pieces with code_decompose.subsplit_oversize; this emitter just removes
    the XLSX *authoring* round-trip and gives the cleaner JSON input the Core supports.)

    Column layout mirrors build_compact exactly: FW_Seq row = [sheet, FW_Reuse,
    <extra flags>, verb@col>=4]; FW_SheetNames = [sheet, prefix|FW_EMPTY_STRING,
    ending|FW_EMPTY_STRING]; data sheet = one row per value (raw=verbatim)."""
    _gate_seq_graph(spec)                                   # STEP 37: reject cycle/missing-operand BEFORE the workbook
    sheets: dict[str, list] = {}

    seq_rows: list[list] = []
    for s in spec.slots:
        cells = {1: s.sheet, 2: "FW_Reuse"}
        col = 3
        for f in s.flags:
            if f != "FW_Reuse":
                cells[col] = f
                col += 1
        vcol = max(col, 4)
        cells[vcol] = cell_for_xlsx(s.verb)                 # verb in col D+ (0-idx 3+)
        mcol = vcol + 1
        if s.separator:
            cells[mcol] = f"FW_Separator({s.separator})"; mcol += 1
        if s.group_replace:
            cells[mcol] = _group_directive(s.group_replace); mcol += 1
        width = max(cells)
        seq_rows.append([cells.get(i) for i in range(1, width + 1)])
    for extra in spec.seq_extra:                            # brace joiners / multi-verb rows, verbatim
        seq_rows.append([cell_for_xlsx(v) for v in extra])
    sheets["FW_Seq"] = seq_rows

    sheets["FW_SheetNames"] = [
        [s.sheet, s.prefix or "FW_EMPTY_STRING", s.ending or "FW_EMPTY_STRING"]
        for s in spec.slots]
    sheets["FW_RunMeFirstOnce"] = [[spec.runme]]
    sheets["FW_Arguments"] = [[a] for a in spec.args] or [["noargs"]]
    sheets["FW_CUSTOM_VAR"] = [[int(c.code), f"FWCUSTOMVAR={c.code} {c.msg}"]
                              for c in spec.custom_vars]
    for s in spec.slots:                                    # data sheets in slot order
        sheets[s.sheet] = [[v if s.raw else cell_for_xlsx(f" {s.key}={v}")] for v in s.values]

    return {"sheets": sheets}


def write_core_json(spec: Spec, path: str | Path) -> Path:
    """Serialise build_compact_core_json(spec) to <path> (UTF-8, no ASCII escaping so
    code stays readable). Returns the Path. Feed to the Core via excel.file=<path>
    (+ core.input.format=json)."""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(build_compact_core_json(spec), ensure_ascii=False, indent=1),
                 encoding="utf-8")
    return p


def build_materialized(spec: Spec, rows: list[str], data_sheet: str = "CANDIDATE",
                       fw_info: str = "dup", clone_path: Optional[str] = None,
                       autofit: float = 0.0) -> Workbook:
    """Pre-expanded workbook: one CANDIDATE sheet holding already-assembled rows,
    emitted by the Core via FW_Combi(1). Used for ablation / N-wise reduced /
    explicit-matrix output (rows are a curated subset of the product)."""
    wb = Workbook()
    seq = wb.active
    seq.title = "FW_Seq"
    seq.cell(1, 1, data_sheet)
    seq.cell(1, 2, "FW_Reuse")
    seq.cell(1, 4, "FW_Combi(1)")
    _write_control_sheets(wb, spec, [data_sheet], seq, fw_info, clone_path)
    cand = wb.create_sheet(data_sheet)
    for r, line in enumerate(rows, start=1):
        cand.cell(r, 1, cell_for_xlsx(line))
    autofit_workbook(wb, autofit)
    return wb


# ---- validation (re-open and check against the verified contract) ----------

CONTROL_HEAD = ["FW_Seq", "FW_SheetNames", "FW_RunMeFirstOnce", "FW_Arguments", "FW_CUSTOM_VAR"]


def validate_workbook(path: str | Path) -> list[str]:
    """Return a list of contract violations ([] == valid)."""
    wb = openpyxl.load_workbook(path, data_only=False)
    sn = wb.sheetnames
    errs: list[str] = []
    if sn[:5] != CONTROL_HEAD:
        errs.append(f"control-sheet head is {sn[:5]}, expected {CONTROL_HEAD}")
    data_sheets = [s for s in sn if not s.startswith("FW_")]
    if not data_sheets:
        errs.append("no data sheets")
    # FW_Seq: each row must carry a recognized Core verb/joiner; its target +
    # any referenced operand sheets must be real data sheets. Accepts the WHOLE
    # Core grammar (Combi/CombiR/Permut[multiset]/PermutR/Subsets+modes/Cartes/
    # Separator/Group/brace), not just FW_Combi.
    for row in wb["FW_Seq"].iter_rows(values_only=True):
        cells = [str(c) for c in row if c not in (None, "")]
        if not cells:
            continue
        target = cells[0]
        headless = is_core_verb(target) or is_core_flag(target)   # row may lead with a directive
        directives = cells if headless else cells[1:]
        if not headless and target not in data_sheets:
            errs.append(f"FW_Seq target '{target}' has no data sheet")
        if not any(is_core_verb(c) for c in directives):
            errs.append(f"FW_Seq row '{target}' has no recognized Core verb/joiner")
        for c in directives:
            for op in verb_sheet_operands(c):
                if op not in data_sheets:
                    errs.append(f"FW_Seq operand sheet '{op}' (in {c!r}) has no data sheet")
    # FW_CUSTOM_VAR col0 int
    for row in wb["FW_CUSTOM_VAR"].iter_rows(values_only=True):
        if row and row[0] is not None:
            try:
                int(row[0])
            except (TypeError, ValueError):
                errs.append(f"FW_CUSTOM_VAR code '{row[0]}' is not an int")
    # data cells must not start with FW_
    for ds in data_sheets:
        for row in wb[ds].iter_rows(values_only=True):
            for v in row:
                if v is not None and str(v).lstrip().startswith("FW_"):
                    errs.append(f"data cell in '{ds}' starts with FW_: {v!r}")
    wb.close()
    return errs


# ---- XLSX -> JSON (lossless round-trip; from v25 combiner_xlsx) -------------

def workbook_to_json(path: str | Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=False)
    return {"workbook": Path(path).name,
            "sheets": [{"name": n,
                        "rows": [list(r) for r in wb[n].iter_rows(values_only=True)]}
                       for n in wb.sheetnames]}


def write_json_sibling(xlsx_path: str | Path) -> Path:
    xlsx_path = Path(xlsx_path)
    out = xlsx_path.with_suffix(".json")
    out.write_text(json.dumps(workbook_to_json(xlsx_path), ensure_ascii=False, indent=2),
                   encoding="utf-8")
    return out


# ===========================================================================
# 5b. Reader->Executor HANDSHAKE awareness  (manage, at GENERATION time, the
#     params the Reader/Executor need but don't reliably emit)
# ---------------------------------------------------------------------------
# The Reader provisions a file-watch handshake the Executor consumes:
#   resultsDbURL/resultsDbURL.properties . sqlTemplate/insert.sql .
#   arguments/{args,fwVar.shift} . runFirstOnce/runmefirstonce.first
# Two of these are routinely WRONG straight out of the Reader (confirmed against
# the live pipeline): fwVar.shift is emitted EMPTY (-> 0 -> the positional FW_VAR
# encoding underflows/overflows the INSERT), and nothing helps you pick a verdict
# code that fits the Results table. fwgen already knows the data-sheet layout at
# generation time, so it can PREDICT the Results columns, VALIDATE the verdict
# code, and EMIT a correct handshake up-front. The Reader's own insert.sql is
# still authoritative at run time; this prediction matches it by construction
# (a projection of the created table) and is used for pre-flight + a ready handshake.
# ===========================================================================

# Base columns the Reader's Results table always carries, in bind order
# (cf. Executor MainWatch.SqlRecord.bind: status, attachment, fwVarOrCustom,
# then combi_id_final/optional/fw_optJ from the candidate filename).
RESULTS_BASE_COLUMNS = ["status", "attachment", "fw_var",
                        "combi_id_final", "combi_id_optional", "fw_optJ"]
# A failing FW_VAR=k sets the (k - shift)-th combos column; shift=1 keeps k>=1 in range.
RECOMMENDED_FWVAR_SHIFT = 1


def results_combos_columns(spec: "Spec") -> list[str]:
    """The boolean combos<i>_<sheet> columns the Reader creates in the Results table,
    numbered 1-based in slot order. Only FW_Exclude/FW_Heading slots are dropped from
    fw_final; FW_Optional slots DO remain as columns (verified against a live Core run:
    fw_final keeps combos<i>_<OPT> columns — the optional dimension is routed to fw_opt<i>
    for the COMBINATORICS but its column still exists)."""
    cols, idx = [], 0
    for s in spec.slots:
        if any(f in s.flags for f in ("FW_Exclude", "FW_Heading")):
            continue
        idx += 1
        cols.append(f"combos{idx}_{s.sheet}")
    return cols


def predict_results_columns(spec: "Spec") -> list[str]:
    """Full Results-table column list the Reader will create (base + combos)."""
    return RESULTS_BASE_COLUMNS + results_combos_columns(spec)


def predict_insert_sql(spec: "Spec", table_name: str) -> str:
    """The INSERT template the Reader hands the Executor (a projection of the
    created table). '?' count == len(columns); == 6 => Executor FW_CUSTOM_VAR mode,
    else FW_VAR mode."""
    cols = predict_results_columns(spec)
    collist = ",\n ".join(f'"{c}"' for c in cols)
    qs = ", ".join("?" for _ in cols)
    return f'INSERT INTO public."{table_name}" ( {collist}\n) VALUES ({qs})'


def handshake_report(spec: "Spec") -> dict:
    """Pre-flight: predicted Results columns, Executor verdict mode, recommended
    fwVar.shift, and verdict-code-overflow warnings."""
    cols = predict_results_columns(spec)
    ncombos = len(results_combos_columns(spec))
    warnings = []
    for cv in spec.custom_vars:
        if cv.code > ncombos:
            warnings.append(
                f"verdict code {cv.code} > #combos columns {ncombos}: with fwVar.shift="
                f"{RECOMMENDED_FWVAR_SHIFT} the positional FW_VAR encoding overflows the INSERT "
                f"(column index out of range). Use a code <= {ncombos}, or add mandatory data sheets.")
    return {"columns": cols, "n_columns": len(cols), "n_combos": ncombos,
            "mode": "FW_CUSTOM_VAR" if len(cols) == 6 else "FW_VAR",
            "fwvar_shift": RECOMMENDED_FWVAR_SHIFT, "warnings": warnings}


def emit_handshake(spec: "Spec", out_dir: str | Path, *, db_url: str,
                   table_name: str, fwvar_shift: int = RECOMMENDED_FWVAR_SHIFT) -> dict:
    """Write a ready-to-run Reader->Executor handshake under out_dir, in the Reader's
    own layout, with the bits the Reader gets wrong fixed up-front: a NON-EMPTY
    fwVar.shift and a verdict-code-safe insert.sql. Returns a manifest (the report +
    files written). Point the Executor's -dir* flags at these subdirs."""
    out = Path(out_dir)
    files = {
        "resultsDbURL/resultsDbURL.properties": db_url.strip() + "\n",
        "sqlTemplate/insert.sql": predict_insert_sql(spec, table_name) + "\n",
        "arguments/args": (" ".join(spec.args)) + "\n",
        "arguments/fwVar.shift": str(int(fwvar_shift)),
        "runFirstOnce/runmefirstonce.first": spec.runme + "\n",
    }
    for rel, content in files.items():
        p = out / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(content, encoding="utf-8")
    (out / "jars").mkdir(parents=True, exist_ok=True)
    rep = handshake_report(spec)
    rep.update(out_dir=str(out), written=sorted(files) + ["jars/"])
    return rep


# ===========================================================================
# 6. Chunking + parallelism  (ported from v25 combiner_xlsx/combiner_chunker)
# ---------------------------------------------------------------------------
# v25 split a huge materialised expansion into chunk_*.txt files and built one
# workbook per chunk, in a ProcessPoolExecutor (openpyxl is CPU-bound, so
# PROCESSES beat the GIL). Here a "chunk" is a slice of the pre-assembled rows;
# each chunk becomes one materialized CANDIDATE workbook on a distinct path
# (no write contention). The Spec dataclass is picklable, so workers rebuild
# their own workbook from (spec, rows_slice).
# ===========================================================================

def _chunk_job(job: tuple) -> str:
    """Top-level (picklable) worker: build + save one chunk workbook, return path."""
    spec, rows, path, data_sheet, fw_info, clone_path, emit_json, autofit = job
    build_materialized(spec, rows, data_sheet=data_sheet, fw_info=fw_info,
                       clone_path=clone_path, autofit=autofit).save(path)
    if emit_json:
        write_json_sibling(path)
    return path


def write_chunked(spec: Spec, rows: list[str], out_dir: str | Path, base: str,
                  chunk_rows: int = 0, fw_info: str = "dup",
                  clone_path: Optional[str] = None, workers: int = 1,
                  emit_json: bool = False, data_sheet: str = "CANDIDATE",
                  autofit: float = 0.0) -> list[str]:
    """Write `rows` as one or more materialized workbooks.

    chunk_rows <= 0  -> a single `<base>.xlsx`.
    chunk_rows  > 0  -> slices of <=chunk_rows rows, each `<base>_chunkNNN.xlsx`.
    workers     > 1  -> build the chunks in a ProcessPoolExecutor.
    Returns the list of written paths."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    size = chunk_rows if chunk_rows and chunk_rows > 0 else (len(rows) or 1)
    slices = [rows[i:i + size] for i in range(0, len(rows), size)] or [[]]
    multi = len(slices) > 1
    jobs = [(spec, sl,
             str(out_dir / (f"{base}_chunk{idx:03d}.xlsx" if multi else f"{base}.xlsx")),
             data_sheet, fw_info, clone_path, emit_json, autofit)
            for idx, sl in enumerate(slices)]

    if workers > 1 and len(jobs) > 1:
        with ProcessPoolExecutor(max_workers=workers) as ex:
            paths = list(ex.map(_chunk_job, jobs))
        log.info("write_chunked: %d chunks via %d processes", len(paths), workers)
    else:
        paths = [_chunk_job(j) for j in jobs]
        log.info("write_chunked: %d chunk(s) sequential", len(paths))
    return paths


# ===========================================================================
# 7. Batch XLSX -> JSON directory  (parallel; from v25 combiner_xlsx)
# ===========================================================================

def _json_job(job: tuple) -> str:
    in_path, out_path = job
    Path(out_path).write_text(
        json.dumps(workbook_to_json(in_path), ensure_ascii=False, indent=2),
        encoding="utf-8")
    return out_path


def xlsx_dir_to_json_dir(xlsx_dir: str | Path, json_dir: Optional[str | Path] = None,
                         workers: int = 1) -> Path:
    """Serialise every `*.xlsx` in `xlsx_dir` to `*.json`. Default `json_dir` is a
    sibling: trailing `_xlsx`->`_json`, else `<name>_json`. Process-parallel
    (CPU-bound load+dump; distinct output paths => no contention)."""
    xlsx_dir = Path(xlsx_dir)
    if json_dir is None:
        nm = (xlsx_dir.name[:-5] + "_json") if xlsx_dir.name.endswith("_xlsx") else xlsx_dir.name + "_json"
        json_dir = xlsx_dir.parent / nm
    json_dir = Path(json_dir)
    json_dir.mkdir(parents=True, exist_ok=True)
    jobs = [(str(b), str(json_dir / (b.stem + ".json")))
            for b in sorted(xlsx_dir.glob("*.xlsx"))]
    if not jobs:
        log.info("xlsx_dir_to_json_dir: no .xlsx in %s", xlsx_dir)
        return json_dir
    if workers > 1 and len(jobs) > 1:
        with ProcessPoolExecutor(max_workers=workers) as ex:
            list(ex.map(_json_job, jobs))
    else:
        for j in jobs:
            _json_job(j)
    log.info("xlsx_dir_to_json_dir: %d workbooks -> %s", len(jobs), json_dir)
    return json_dir


# ===========================================================================
# 8. Domain-facing PREVIEW + spec AUTHORING (landing / builder helpers)
# ---------------------------------------------------------------------------
# These speak the specialist's language and never expose FW_ machinery. The
# preview renders an example; the authoring helpers turn a specialist's listed
# parameters (slots + values) into a spec — shared by the CLI wizard/scaffold
# and the GUI Builder. Everything else (baseline, goal direction, FW_*,
# validation) is auto-derived by parse_spec.
# ===========================================================================

def _combo_at(spec: Spec, index: int) -> tuple[str, ...]:
    """The `index`-th combination in itertools.product order (rightmost slot
    fastest) WITHOUT materialising the product. index 0 == baseline."""
    out = [""] * len(spec.slots)
    for i in range(len(spec.slots) - 1, -1, -1):
        vals = spec.slots[i].values
        out[i] = vals[index % len(vals)]
        index //= len(vals)
    return tuple(out)


def sample_combos(spec: Spec, k: int = 6) -> list[tuple[str, ...]]:
    """baseline + up to k-1 evenly-spread samples (no full materialisation)."""
    total = spec.combos
    if total <= k:
        return [_combo_at(spec, i) for i in range(total)]
    idxs = [0] + [round(i * (total - 1) / (k - 1)) for i in range(1, k)]
    seen, out = set(), []
    for ix in idxs:
        if ix not in seen:
            seen.add(ix)
            out.append(_combo_at(spec, ix))
    return out


def preview_line(spec: Spec, combo) -> str:
    """A candidate in the specialist's language (no FW_)."""
    return "  ·  ".join(f"{s.key}={v}" for s, v in zip(spec.slots, combo))


def render_preview(spec: Spec, k: int = 6) -> str:
    L = [f"━━ {spec.title} ━━"]
    if spec.note:
        L.append(spec.note)
    L += ["", f"Parameters (slots) — {len(spec.slots)}; full space = {spec.combos} candidates:"]
    for s in spec.slots:
        L.append(f"  • {s.sheet:<12} : " + "  |  ".join(s.values))
    if spec.goals:
        L += ["", "Optimize (Pareto): " + spec.goals_property()]
    samples = sample_combos(spec, k)
    L += ["", f"Example candidates ({len(samples)} of {spec.combos}):"]
    for i, c in enumerate(samples):
        tag = "   ← baseline" if i == 0 and c == spec.baseline else ""
        L.append(f"  {i+1}. {preview_line(spec, c)}{tag}")
    return "\n".join(L)


def render_preview_html(specs: list[Spec]) -> str:
    import html
    p = ["<!doctype html><meta charset=utf-8><title>fwgen — examples</title>",
         "<style>body{font:14px/1.6 system-ui,sans-serif;max-width:920px;margin:2rem auto;padding:0 1rem;color:#222}"
         "h2{border-bottom:2px solid #ddd;margin-top:2rem}code{background:#f4f4f4;padding:1px 5px;border-radius:3px}"
         ".cand{font-family:monospace;background:#fafafa;border-left:3px solid #4a90d9;padding:.25rem .6rem;margin:.2rem 0}"
         ".slot b{display:inline-block;min-width:9rem;color:#1a4}</style>",
         "<h1>fwgen — example input scenarios</h1>",
         "<p>Each card is a domain search space. To make your own: list <b>your</b> slots and values — "
         "baseline, goals and file format are derived automatically.</p>"]
    for sp in specs:
        p.append(f"<h2>{html.escape(sp.title)}</h2>")
        if sp.note:
            p.append(f"<p>{html.escape(sp.note)}</p>")
        p.append(f"<p><b>{len(sp.slots)} parameters</b>, full space = <b>{sp.combos}</b> candidates &middot; "
                 f"optimize <code>{html.escape(sp.goals_property() or 'n/a')}</code></p><div class=slot>")
        for s in sp.slots:
            p.append(f"<div><b>{html.escape(s.sheet)}</b> {html.escape('  |  '.join(s.values))}</div>")
        p.append("</div>")
        for c in sample_combos(sp, 5):
            p.append(f"<div class=cand>{html.escape(preview_line(sp, c))}</div>")
    return "\n".join(p)


# ---- authoring: a specialist's listed parameters -> spec --------------------

def build_spec_dict(name, slots, *, title=None, goals=None, args=None,
                    note="", custom_vars=None) -> dict:
    """`slots`: list of (sheet, key_or_None, values_list). Returns a raw spec
    dict (same shape as a spec file); pass to parse_spec() to validate."""
    d = {"title": title or name}
    if note:
        d["note"] = note
    if goals:
        d["goals"] = list(goals)
    if args:
        d["args"] = list(args)
    d["slots"] = [({"sheet": s, "values": list(vals)} | ({"key": k} if k else {}))
                  for (s, k, vals) in slots]
    if custom_vars:
        d["custom_vars"] = [{"code": int(c), "msg": m} for (c, m) in custom_vars]
    return d


def _toml_basic(s) -> str:
    return '"' + str(s).replace("\\", "\\\\").replace('"', '\\"') + '"'


def dump_spec_toml(data: dict, *, comments: bool = False) -> str:
    """Serialise a raw spec dict to TOML (round-trips through parse_spec).
    `comments=True` adds beginner hints (scaffold-template mode)."""
    L = []
    if data.get("title"):
        L.append(f'title = {_toml_basic(data["title"])}')
    if data.get("note"):
        L.append(f'note  = {_toml_basic(data["note"])}')
    goals = data.get("goals")
    if goals and all(isinstance(g, str) for g in goals):
        L.append("goals = [" + ", ".join(_toml_basic(g) for g in goals) + "]"
                 + ("   # direction auto-inferred (cost/latency→min, f1/acc→max)" if comments else ""))
    if data.get("args"):
        L.append("args  = [" + ", ".join(_toml_basic(a) for a in data["args"]) + "]")
    if comments:
        L += ["", "# Each [[slots]] is one parameter axis the engine combines.",
              "# Row 0 of `values` = baseline. Sheet name must NOT start with FW_."]
    for s in data.get("slots", []):
        L += ["", "[[slots]]", f'sheet  = {_toml_basic(s["sheet"])}']
        if s.get("key"):
            L.append(f'key    = {_toml_basic(s["key"])}')
        L.append("values = [" + ", ".join(_toml_basic(v) for v in s["values"]) + "]")
    if goals and not all(isinstance(g, str) for g in goals):     # explicit [[goals]]
        for g in goals:
            L += ["", "[[goals]]", f'key = {_toml_basic(g["key"])}',
                  f'dir = {_toml_basic(g.get("dir") or g.get("direction") or "max")}']
    for c in data.get("custom_vars", []):
        L += ["", "[[custom_vars]]", f'code = {int(c["code"])}', f'msg  = {_toml_basic(c["msg"])}']
    return "\n".join(L) + "\n"


def scaffold_template(name: str, from_spec: Optional[Spec] = None) -> str:
    """A commented TOML template a specialist fills in. With `from_spec`, its
    slots/goals are pre-filled as a worked example to edit."""
    if from_spec:
        data = {"title": f"{name} (from example: {from_spec.name})",
                "goals": [g.key for g in from_spec.goals],
                "slots": [{"sheet": s.sheet, "key": s.key, "values": list(s.values)}
                          for s in from_spec.slots]}
    else:
        data = {"title": name, "goals": ["accuracy", "cost"],
                "slots": [{"sheet": "PARAM1", "key": "param1",
                           "values": ["baseline", "variant_a", "variant_b"]},
                          {"sheet": "PARAM2", "key": "param2", "values": ["x", "y"]}]}
    header = ("# fwgen spec template — replace with YOUR parameters, then:\n"
              f"#   python3 fwgen_cli.py preview --specs <dir> --spec {name}\n"
              "#   python3 fwgen_cli.py gen     --specs <dir> --out out\n"
              "# Only slots+values are required; baseline = first value; goals optional.\n\n")
    return header + dump_spec_toml(data, comments=True)


# ===========================================================================
# 9. Verb-coverage SUITE  (isolated per-verb chunks — NO cartesian explosion)
# ---------------------------------------------------------------------------
# To prove the generator can emit data for EVERY combination rule the Core runs,
# the right shape is MANY tiny single-purpose workbooks (one verb each, ≤2 slots,
# ≤3 values) — NOT one workbook listing all verbs as slots (whose product would be
# a "cosmic" cartesian). Each workbook here exercises one verb's hot path on its
# own; the Core's per-workbook output stays small (≤ a few dozen rows). Braces are
# covered separately (see test_fwgen.test_brace_joiner_seq_extra) because they need
# excluded operand sheets.
# ===========================================================================

COVERAGE_MAX_SLOTS = 2     # an isolated chunk is 1 slot (2 for the binary Cartes)
COVERAGE_MAX_VALUES = 3    # tiny value sets keep every chunk's product bounded


def verb_coverage_specs() -> list[Spec]:
    """Ready-to-build Specs, one per Core combinatorial verb form, each isolated and
    tiny (bounded by COVERAGE_MAX_SLOTS / COVERAGE_MAX_VALUES). Covers Combi(1|k|all),
    CombiR(k), Permut / Permut(k) / multiset, PermutR(k), Subsets + all five modes,
    and binary Cartes."""
    def _slot(sheet, verb, vals, key=None, flags=None):
        d = {"sheet": sheet, "values": vals, "verb": verb}
        if key:
            d["key"] = key
        if flags:
            d["flags"] = flags
        return d

    def _spec(name, slots):
        return parse_spec({"title": f"verb coverage: {name}",
                           "args": ["mode=dsl_coverage"],
                           "custom_vars": [{"code": 2, "msg": "FW_Seq directive rejected / parser error"}],
                           "slots": slots}, name)

    return [
        _spec("cov_combi1",          [_slot("LEAF",  "FW_Combi(1)", ["x", "y"])]),
        _spec("cov_combi2",          [_slot("PAIR",  "FW_Combi(2)", ["p", "q", "r"])]),
        _spec("cov_combi_all",       [_slot("ANY",   "FW_Combi(all)", ["p", "q", "r"])]),
        _spec("cov_combiR2",         [_slot("MULTI", "FW_CombiR(2)", ["a", "b"])]),
        _spec("cov_permut",          [_slot("ORDER", "FW_Permut", ["s1", "s2", "s3"])]),
        _spec("cov_permut_k",        [_slot("KPERM", "FW_Permut(2)", ["s1", "s2", "s3"])]),
        _spec("cov_permut_multiset", [_slot("MSET",  "FW_Permut(PermutationGenerator.TreatDuplicatesAs.IDENTICAL)", ["d", "d", "e"])]),
        _spec("cov_permutR2",        [_slot("REP",   "FW_PermutR(2)", ["0", "1"])]),
        _spec("cov_subsets",         [_slot("POW",   "FW_Subsets", ["h1", "h2", "h3"])]),
        _spec("cov_subsets_exact2",  [_slot("EX",    "FW_Subsets_EXACT(2)", ["m", "n", "o"])]),
        _spec("cov_subsets_range",   [_slot("RNG",   "FW_Subsets_RANGE(1,2)", ["m", "n", "o"])]),
        _spec("cov_subsets_before",  [_slot("BEF",   "FW_Subsets_BEFORE(2)", ["m", "n", "o"])]),
        _spec("cov_subsets_after",   [_slot("AFT",   "FW_Subsets_AFTER(1)", ["m", "n", "o"])]),
        _spec("cov_subsets_given",   [_slot("GIV",   "FW_Subsets_GIVEN(1,3)", ["m", "n", "o"])]),
        _spec("cov_cartes",          [_slot("CARTA", "FW_Cartes(CARTB)", ["u", "v"]),
                                      _slot("CARTB", "FW_Combi(1)", ["w"])]),
    ]


INTEROP_MAX_CORE_ROWS = 8000   # the 10+-row interop workbook must stay below this


def verb_interop_spec() -> Spec:
    """ONE workbook with 10+ FW_Seq rows of DIFFERENT verbs interoperating — to cover
    multi-row interoperability (a realistic 'long' FW_Seq), NOT just isolated verbs.
    Values are deliberately tiny so the inner combined data — the full product the
    Core builds — stays small (estimate_core_combos ≈ 2304, well under
    INTEROP_MAX_CORE_ROWS). HEAD/FOOT are constant framing slots (cf. testgen_api)."""
    def _slot(sheet, verb, vals):
        return {"sheet": sheet, "values": vals, "verb": verb}
    slots = [
        _slot("HEAD",    "FW_Combi(1)",          ["prelude"]),            # 1 (framing constant)
        _slot("PROFILE", "FW_Combi(1)",          ["free", "pro"]),        # 2  cartesian leaf
        _slot("REGION",  "FW_Combi(1)",          ["us", "eu"]),           # 2  cartesian leaf
        _slot("TIER",    "FW_Combi(1)",          ["std", "prem"]),        # 2  cartesian leaf
        _slot("MODE",    "FW_Combi(1)",          ["sync", "async"]),      # 2  cartesian leaf
        _slot("PAIR",    "FW_Combi(2)",          ["alpha", "beta", "gamma"]),  # C(3,2)=3
        _slot("ORDER",   "FW_Permut",            ["init", "run"]),        # 2! = 2  (ordering bugs)
        _slot("HDR",     "FW_Subsets",           ["gzip", "auth"]),       # 2^2 = 4 (optional headers)
        _slot("PICK",    "FW_Subsets_EXACT(1)",  ["v1", "v2"]),           # C(2,1)=2
        _slot("COMBO",   "FW_CombiR(2)",         ["e1", "e2"]),           # multicombi = 3
        _slot("FOOT",    "FW_Combi(1)",          ["coda"]),               # 1 (framing constant)
    ]
    return parse_spec({"title": "10+ row FW_Seq interop (mixed verbs; bounded inner product)",
                       "args": ["mode=dsl_interop"],
                       "custom_vars": [{"code": 2, "msg": "FW_Seq directive rejected / parser error"}],
                       "slots": slots}, "cov_interop_11rows")


def build_verb_coverage(out_dir: str | Path, fw_info: str = "skip",
                        autofit: float = 0.0, emit_json: bool = False) -> list[dict]:
    """Emit the coverage suite into `out_dir`: one tiny ISOLATED workbook per Core
    verb (no cartesian-of-verbs explosion) PLUS one 10+-row mixed-verb interop
    workbook (bounded inner product). Returns a manifest list
    [{name, verbs, fw_seq_rows, est_core_rows, valid, path}]."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest = []
    for spec in verb_coverage_specs() + [verb_interop_spec()]:
        path = out_dir / f"{spec.name}.xlsx"
        build_compact(spec, fw_info=fw_info, autofit=autofit).save(path)
        if emit_json:
            write_json_sibling(path)
        errs = validate_workbook(path)
        manifest.append({"name": spec.name,
                         "verbs": " | ".join(dict.fromkeys(s.verb for s in spec.slots)),
                         "fw_seq_rows": len(spec.slots) + len(spec.seq_extra),
                         "est_core_rows": estimate_core_combos(spec),
                         "valid": "OK" if not errs else "; ".join(errs),
                         "path": str(path)})
    return manifest
